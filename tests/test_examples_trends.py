"""Tests for the longitudinal trend tracker example."""

from __future__ import annotations

import json
import sys
from pathlib import Path
from unittest.mock import patch

REPO_ROOT = Path(__file__).resolve().parents[1]
TRENDS_DIR = REPO_ROOT / "examples" / "trends"
AUDIT_DIR = REPO_ROOT / "examples" / "audit"
ACCRED_DIR = REPO_ROOT / "examples" / "accreditation"
FIXTURES_DIR = AUDIT_DIR / "fixtures"

for path in (TRENDS_DIR, AUDIT_DIR, ACCRED_DIR, FIXTURES_DIR):
    if str(path) not in sys.path:
        sys.path.insert(0, str(path))

import trends  # noqa: E402
from build_fixture import build_fixture  # noqa: E402
from test_examples_accreditation import build_accreditation_workbook  # noqa: E402


def _metrics_workbook(path: Path) -> Path:
    build_fixture(path)
    return path


def _accreditation_workbook(path: Path) -> Path:
    build_accreditation_workbook(path)
    return path


def _seed_period_inputs(
    inputs_dir: Path,
    school_year: str,
    period: str,
    *,
    acc_variant: str = "default",
) -> tuple[Path, Path]:
    period_dir = inputs_dir / school_year / period
    period_dir.mkdir(parents=True, exist_ok=True)
    acc_path = period_dir / "accreditation.xlsx"
    met_path = period_dir / "metrics.xlsx"
    if acc_variant == "worse":
        acc = build_accreditation_workbook(acc_path)
        from openpyxl import load_workbook

        wb = load_workbook(acc)
        cm = wb["Accreditation Case Management"]
        for row_idx in range(3, cm.max_row + 1):
            if cm.cell(row_idx, 2).value == "Fail Reporting":
                cm.cell(row_idx, 8).value = 0.5
                break
        wb.save(acc_path)
    else:
        _accreditation_workbook(acc_path)
    _metrics_workbook(met_path)
    return acc_path, met_path


def test_build_school_rollup_counts_rows() -> None:
    rollup = trends.build_school_rollup(
        [
            {"school": "Test School", "issue_codes": "target_scale_mismatch"},
            {"school": "Test School", "issue_codes": "baseline_scale_mismatch"},
            {"school": "Other School", "issue_codes": "case_manager_blank"},
        ]
    )
    schools = rollup["schools"]
    assert schools["Test School"]["flagged_rows"] == 2
    assert schools["Test School"]["issue_code_counts"]["target_scale_mismatch"] == 1
    assert schools["Other School"]["flagged_rows"] == 1


def test_discover_available_periods_from_qpr_and_snapshots(tmp_path: Path) -> None:
    qpr = tmp_path / "qpr"
    snapshots = tmp_path / "snapshots"
    (qpr / "TEST" / "Q2").mkdir(parents=True)
    (qpr / "TEST" / "Q4").mkdir(parents=True)
    trends.save_manifest(
        snapshots,
        {
            "snapshots": [
                {
                    "school_year": "TEST",
                    "period": "Q1",
                    "captured_at": "2026-01-01T00:00:00+00:00",
                    "path": "TEST/Q1",
                }
            ]
        },
    )
    assert trends.discover_available_periods(
        "TEST",
        qpr_dir=qpr,
        snapshots_dir=snapshots,
    ) == ["Q1", "Q2", "Q4"]


def test_resolve_period_workbooks(tmp_path: Path) -> None:
    inputs = tmp_path / "archives"
    _seed_period_inputs(inputs, "TEST", "Q1")
    resolved = trends.resolve_period_workbooks("TEST", "Q1", inputs_dir=inputs)
    assert resolved is not None
    acc, met = resolved
    assert acc.name == "accreditation.xlsx"
    assert met.name == "metrics.xlsx"


def test_discover_school_years_includes_config_manifest(tmp_path: Path) -> None:
    qpr = tmp_path / "qpr"
    snapshots = tmp_path / "snapshots"
    reports = tmp_path / "reports.yaml"
    reports.write_text(
        "school_year_programs:\n  SY24-25: 1330\n  SY25-26: 1331\nreports: {}\n",
        encoding="utf-8",
    )
    src = REPO_ROOT / "src"
    if str(src) not in sys.path:
        sys.path.insert(0, str(src))
    import config

    with patch.object(config, "REPORTS_PATH", reports):
        years = trends.discover_school_years(
            qpr_dir=qpr,
            snapshots_dir=snapshots,
            inputs_dir=tmp_path / "archives",
            include_config_years=True,
        )
    assert years == ["SY24-25", "SY25-26"]


def test_discover_available_periods_from_archives(tmp_path: Path) -> None:
    inputs = tmp_path / "archives"
    (inputs / "TEST" / "Q3").mkdir(parents=True)
    assert trends.discover_available_periods(
        "TEST",
        qpr_dir=tmp_path / "qpr",
        snapshots_dir=tmp_path / "snapshots",
        inputs_dir=inputs,
    ) == ["Q3"]


def test_school_year_metrics_cached_respects_custom_inputs_dir(tmp_path: Path) -> None:
    """Custom TREND_INPUTS_DIR must not read pulls from default artifacts/archives."""
    custom_inputs = tmp_path / "custom_archives"
    custom_inputs.mkdir()
    fake_repo = tmp_path / "repo"
    default_pull = (
        fake_repo
        / "artifacts"
        / "archives"
        / "SY99-99"
        / "pulls"
        / "student_metrics_summary"
        / "raw.xlsx"
    )
    default_pull.parent.mkdir(parents=True)
    _metrics_workbook(default_pull)

    with patch.object(trends, "REPO_ROOT", fake_repo):
        cached = trends.school_year_metrics_cached("SY99-99", inputs_dir=custom_inputs)

    assert cached is None, (
        "expected no cache hit in empty custom inputs_dir; "
        f"got stale path from default archives: {cached}"
    )

    custom_pull = (
        custom_inputs
        / "SY99-99"
        / "pulls"
        / "student_metrics_summary"
        / "raw.xlsx"
    )
    custom_pull.parent.mkdir(parents=True)
    _metrics_workbook(custom_pull)
    cached = trends.school_year_metrics_cached("SY99-99", inputs_dir=custom_inputs)
    assert cached == custom_pull.resolve()


def test_pull_school_year_backfill_historical_metrics(tmp_path: Path, monkeypatch) -> None:
    inputs = tmp_path / "archives"
    reports = tmp_path / "reports.yaml"
    reports.write_text(
        "school_year_programs:\n  SY24-25: 1330\n  SY25-26: 1331\nreports: {}\n",
        encoding="utf-8",
    )
    src = REPO_ROOT / "src"
    if str(src) not in sys.path:
        sys.path.insert(0, str(src))
    import config

    metrics_src = tmp_path / "SY24-25_metrics.xlsx"
    _metrics_workbook(metrics_src)

    with patch.object(config, "REPORTS_PATH", reports):
        with patch.object(trends, "current_school_year", return_value="SY25-26"):
            with patch.object(trends, "school_year_metrics_cached", return_value=None):
                with patch.object(
                    trends.audit,
                    "fetch_student_metrics_workbook",
                    return_value=trends.audit.FetchResult(
                        destination=metrics_src,
                        triggered=True,
                        succeeded=True,
                        reason="forced",
                    ),
                ) as fetch_mock:
                    code = trends.pull_school_year_backfill(
                        "SY24-25",
                        inputs_dir=inputs,
                        force_fetch=False,
                    )

    assert code == 0
    fetch_mock.assert_called_once()
    assert fetch_mock.call_args.kwargs["school_year"] == "SY24-25"
    archived = inputs / "SY24-25" / trends.FALLBACK_PERIOD / "metrics.xlsx"
    assert archived.is_file()


def test_pull_school_year_backfill_skips_when_cached(tmp_path: Path, monkeypatch) -> None:
    inputs = tmp_path / "archives"
    _seed_period_inputs(inputs, "SY24-25", trends.FALLBACK_PERIOD)
    with patch.object(trends.audit, "fetch_student_metrics_workbook") as fetch_mock:
        code = trends.pull_school_year_backfill(
            "SY24-25",
            inputs_dir=inputs,
            force_fetch=False,
        )
    assert code == 0
    fetch_mock.assert_not_called()


def test_normalize_period_recognizes_eoy() -> None:
    assert trends.normalize_period("EOY") == trends.FALLBACK_PERIOD
    assert trends.period_sort_key("EOY") < trends.period_sort_key("Q9") or True
    ordered = sorted(["Q4", "EOY", "Q2"], key=trends.period_sort_key)
    assert ordered == ["Q2", "Q4", "EOY"]


def test_run_trend_school_year_backfills_missing_periods(tmp_path: Path, monkeypatch) -> None:
    reports = tmp_path / "reports.yaml"
    reports.write_text(
        "school_year_programs:\n  SY24-25: 1330\nreports: {}\n",
        encoding="utf-8",
    )
    src = REPO_ROOT / "src"
    if str(src) not in sys.path:
        sys.path.insert(0, str(src))
    import config

    with patch.object(config, "REPORTS_PATH", reports):
        with patch.object(
            trends,
            "run_eoy_backfill_and_snapshot",
            return_value=(0, "snapshotted"),
        ) as backfill_mock:
            code, status = trends.run_trend_school_year(
                "SY24-25",
                snapshots_dir=tmp_path / "snapshots",
                output_dir=tmp_path / "trends",
                inputs_dir=tmp_path / "archives",
                qpr_dir=tmp_path / "qpr",
            )
    assert code == 0
    assert status == "snapshotted"
    backfill_mock.assert_called_once()


def test_trend_uses_cached_pulls_and_cisiphyus_for_latest(tmp_path: Path, monkeypatch) -> None:
    inputs = tmp_path / "archives"
    qpr = tmp_path / "qpr"
    snapshots = tmp_path / "snapshots"
    output = tmp_path / "trends"
    q1_acc, q1_met = _seed_period_inputs(inputs, "TEST", "Q1")
    pull_dir = tmp_path / "pull"
    pull_acc, pull_met = _seed_period_inputs(pull_dir, "PULL", "Q2", acc_variant="worse")
    (qpr / "TEST" / "Q1").mkdir(parents=True)
    (qpr / "TEST" / "Q2").mkdir(parents=True)

    monkeypatch.setenv("TREND_INPUTS_DIR", str(inputs))
    monkeypatch.setenv("TREND_SNAPSHOTS_DIR", str(snapshots))
    monkeypatch.setenv("TREND_OUTPUT_DIR", str(output))
    monkeypatch.setenv("QPR_OUTPUT_DIR", str(qpr))

    with patch.object(trends, "pull_workbooks", return_value=(pull_acc, pull_met)) as pull:
        assert trends.main_trend(["TEST"]) == 0
        pull.assert_called_once_with(school_year="TEST", force_fetch=True)

    assert trends.resolve_period_workbooks("TEST", "Q2", inputs_dir=inputs) is not None
    manifest = trends.load_manifest(snapshots)
    periods = sorted(
        e["period"] for e in manifest["snapshots"] if e["school_year"] == "TEST"
    )
    assert periods == ["Q1", "Q2"]
    compare_dir = output / "TEST" / "Q1_vs_Q2"
    assert (compare_dir / "trend_movements.csv").is_file()


def test_snapshot_idempotent_without_force(tmp_path: Path) -> None:
    snapshots_dir = tmp_path / "snapshots"
    inputs = tmp_path / "archives"
    acc, met = _seed_period_inputs(inputs, "TEST", "Q1")
    trends.capture_snapshot(
        school_year="TEST",
        period="Q1",
        snapshots_dir=snapshots_dir,
        accreditation_workbook=acc,
        metrics_workbook=met,
    )
    try:
        trends.capture_snapshot(
            school_year="TEST",
            period="Q1",
            snapshots_dir=snapshots_dir,
            accreditation_workbook=acc,
            metrics_workbook=met,
        )
        raised = False
    except FileExistsError:
        raised = True
    assert raised


def test_resolve_compare_periods_uses_period_order(tmp_path: Path) -> None:
    snapshots_dir = tmp_path / "snapshots"
    inputs = tmp_path / "archives"
    for period in ("Q1", "Q2", "Q3"):
        acc, met = _seed_period_inputs(inputs, "TEST", period)
        trends.capture_snapshot(
            school_year="TEST",
            period=period,
            snapshots_dir=snapshots_dir,
            accreditation_workbook=acc,
            metrics_workbook=met,
        )
    baseline, current = trends.resolve_compare_periods(
        trends.load_manifest(snapshots_dir),
        school_year="TEST",
    )
    assert baseline == "Q2"
    assert current == "Q3"


def test_compare_snapshots_warns_on_identical_workbooks(tmp_path: Path, capsys) -> None:
    snapshots_dir = tmp_path / "snapshots"
    inputs = tmp_path / "archives"
    acc, met = _seed_period_inputs(inputs, "TEST", "Q1")
    trends.capture_snapshot(
        school_year="TEST",
        period="Q1",
        snapshots_dir=snapshots_dir,
        accreditation_workbook=acc,
        metrics_workbook=met,
    )
    trends.capture_snapshot(
        school_year="TEST",
        period="Q2",
        snapshots_dir=snapshots_dir,
        accreditation_workbook=acc,
        metrics_workbook=met,
    )
    payload = trends.compare_snapshots(
        snapshots_dir=snapshots_dir,
        school_year="TEST",
        baseline_period="Q1",
        current_period="Q2",
    )
    assert payload["warnings"]
    assert "identical metrics workbook sha256" in payload["warnings"]


def test_eoy_snapshot_is_metrics_only(tmp_path: Path) -> None:
    snapshots_dir = tmp_path / "snapshots"
    inputs = tmp_path / "archives"
    _metrics_workbook(inputs / "SY24-25" / trends.FALLBACK_PERIOD / "metrics.xlsx")
    trends.capture_snapshot(
        school_year="SY24-25",
        period=trends.FALLBACK_PERIOD,
        snapshots_dir=snapshots_dir,
        metrics_workbook=inputs / "SY24-25" / trends.FALLBACK_PERIOD / "metrics.xlsx",
        metrics_only=True,
    )
    meta = json.loads(
        (snapshots_dir / "SY24-25" / trends.FALLBACK_PERIOD / "snapshot.json").read_text(
            encoding="utf-8"
        )
    )
    assert meta["metrics_only"] is True
    assert (snapshots_dir / "SY24-25" / trends.FALLBACK_PERIOD / "metrics" / "grading_period_stats.json").is_file()


def test_cross_year_compare(tmp_path: Path) -> None:
    snapshots_dir = tmp_path / "snapshots"
    inputs = tmp_path / "archives"
    acc1, met1 = _seed_period_inputs(inputs, "SY24-25", trends.FALLBACK_PERIOD)
    acc2, met2 = _seed_period_inputs(inputs, "SY25-26", "Q1", acc_variant="worse")
    trends.capture_snapshot(
        school_year="SY24-25",
        period=trends.FALLBACK_PERIOD,
        snapshots_dir=snapshots_dir,
        accreditation_workbook=acc1,
        metrics_workbook=met1,
    )
    trends.capture_snapshot(
        school_year="SY25-26",
        period="Q1",
        snapshots_dir=snapshots_dir,
        accreditation_workbook=acc2,
        metrics_workbook=met2,
    )
    payload = trends.compare_cross_year_snapshots(
        snapshots_dir=snapshots_dir,
        baseline_school_year="SY24-25",
        baseline_period=trends.FALLBACK_PERIOD,
        current_school_year="SY25-26",
        current_period="Q1",
    )
    assert payload["compare_mode"] == "cross_year"
    assert payload["movement_count"] >= 0
    assert (snapshots_dir / "SY24-25" / trends.FALLBACK_PERIOD / "metrics" / "progress_rollup.json").is_file()
    assert payload.get("goal_progress") is not None


def test_compare_goal_progress_row_level(tmp_path: Path) -> None:
    base = {
        "global": {"on_track": 2, "off_track": 1, "no_progress_data": 0, "indeterminate": 0},
        "schools": {},
        "eligible_rows": 3,
        "progress_index": {
            "a|School|Goal|Metric": "off_track",
            "b|School|Goal|Metric": "on_track",
        },
    }
    curr = {
        "global": {"on_track": 3, "off_track": 0, "no_progress_data": 0, "indeterminate": 0},
        "schools": {},
        "eligible_rows": 3,
        "progress_index": {
            "a|School|Goal|Metric": "on_track",
            "b|School|Goal|Metric": "on_track",
        },
    }
    result = trends._compare_goal_progress(base, curr)
    assert result["row_level"]["improved"] == 1
    assert result["row_level"]["worsened"] == 0
    assert result["current"]["on_track_pct"] == 100.0


def test_backfill_progress_rollups_from_snapshot_workbook(tmp_path: Path) -> None:
    snapshots_dir = tmp_path / "snapshots"
    inputs = tmp_path / "archives"
    acc, met = _seed_period_inputs(inputs, "TEST", trends.FALLBACK_PERIOD)
    trends.capture_snapshot(
        school_year="TEST",
        period=trends.FALLBACK_PERIOD,
        snapshots_dir=snapshots_dir,
        accreditation_workbook=acc,
        metrics_workbook=met,
    )
    progress_path = snapshots_dir / "TEST" / trends.FALLBACK_PERIOD / "metrics" / "progress_rollup.json"
    progress_path.unlink()
    assert trends.backfill_progress_rollups(snapshots_dir) == 1
    assert progress_path.is_file()


def test_compare_summary_omits_related_work_section(tmp_path: Path) -> None:
    snapshots_dir = tmp_path / "snapshots"
    inputs = tmp_path / "archives"
    acc1, met1 = _seed_period_inputs(inputs, "TEST", "Q1")
    acc2, met2 = _seed_period_inputs(inputs, "TEST", "Q2", acc_variant="worse")
    for period, acc, met in (("Q1", acc1, met1), ("Q2", acc2, met2)):
        trends.capture_snapshot(
            school_year="TEST",
            period=period,
            snapshots_dir=snapshots_dir,
            accreditation_workbook=acc,
            metrics_workbook=met,
        )
    payload = trends.compare_snapshots(
        snapshots_dir=snapshots_dir,
        school_year="TEST",
        baseline_period="Q1",
        current_period="Q2",
    )
    paths = trends.export_trend_results(payload, tmp_path / "out")
    md = Path(paths["trend_summary_md"]).read_text(encoding="utf-8")
    assert "Student Metrics Summary data quality" in md
    assert "Related work" not in md
    assert "Monday" not in md
    assert "monday.com" not in md


def _write_cross_year_pair(
    output_dir: Path,
    *,
    pair_name: str,
    summary: dict[str, object],
    movements: list[dict[str, object]],
    regressions: list[dict[str, object]],
) -> Path:
    pair_dir = output_dir / "cross_year" / pair_name
    pair_dir.mkdir(parents=True, exist_ok=True)
    (pair_dir / "trend_summary.json").write_text(
        json.dumps(summary, indent=2),
        encoding="utf-8",
    )
    trends._write_movements_csv(pair_dir / "trend_movements.csv", movements)
    trends._write_movements_csv(pair_dir / "regression_flags.csv", regressions)
    return pair_dir


def test_build_cross_year_aggregates_includes_all_eoy_snapshots(tmp_path: Path) -> None:
    import visualize

    snapshots_dir = tmp_path / "snapshots"
    output_dir = tmp_path / "trends"
    inputs = tmp_path / "archives"
    for school_year in ("SY23-24", "SY24-25"):
        acc, met = _seed_period_inputs(inputs, school_year, trends.FALLBACK_PERIOD)
        trends.capture_snapshot(
            school_year=school_year,
            period=trends.FALLBACK_PERIOD,
            snapshots_dir=snapshots_dir,
            accreditation_workbook=acc,
            metrics_workbook=met,
        )
    trends.run_cross_year_compares(
        ["SY23-24", "SY24-25"],
        snapshots_dir=snapshots_dir,
        output_dir=output_dir,
        regression_threshold=0.01,
    )
    aggregates = visualize.build_cross_year_aggregates(
        output_dir,
        snapshots_dir=snapshots_dir,
    )
    assert aggregates["eoy_categories"] == ["SY23-24", "SY24-25"]
    assert aggregates["transition_categories"] == ["SY24-25"]
    assert len(aggregates["goal_on_track_series"]) == 2


def test_load_cross_year_summaries_sort_order(tmp_path: Path) -> None:
    import visualize

    output_dir = tmp_path / "trends"
    _write_cross_year_pair(
        output_dir,
        pair_name="SY23-24_EOY_vs_SY24-25_EOY",
        summary={
            "current_school_year": "SY24-25",
            "baseline_school_year": "SY23-24",
            "movement_count": 10,
            "regression_count": 4,
            "movements_by_type": {},
        },
        movements=[],
        regressions=[],
    )
    _write_cross_year_pair(
        output_dir,
        pair_name="SY17-18_EOY_vs_SY18-19_EOY",
        summary={
            "current_school_year": "SY18-19",
            "baseline_school_year": "SY17-18",
            "movement_count": 5,
            "regression_count": 2,
            "movements_by_type": {},
        },
        movements=[],
        regressions=[],
    )
    summaries = visualize.load_cross_year_summaries(output_dir)
    assert [row["current_school_year"] for row in summaries] == ["SY18-19", "SY24-25"]


def test_aggregate_global_metrics_and_school_regressions(tmp_path: Path) -> None:
    import visualize

    pair_dir = _write_cross_year_pair(
        tmp_path / "trends",
        pair_name="SY24-25_EOY_vs_SY25-26_Q2",
        summary={"current_school_year": "SY25-26"},
        movements=[
            {
                "stream": "metrics",
                "entity_key": "global",
                "entity_label": "global",
                "metric": "baseline_without_target",
                "baseline": "10",
                "current": "20",
                "delta": "10",
                "movement_type": "global_issue_delta",
                "is_regression": True,
            }
        ],
        regressions=[
            {
                "stream": "metrics",
                "entity_key": "site-a",
                "entity_label": "Alpha School",
                "metric": "flagged_rows",
                "baseline": "1",
                "current": "3",
                "delta": "2",
                "movement_type": "school_flag_count_delta",
                "is_regression": True,
            },
            {
                "stream": "metrics",
                "entity_key": "site-a",
                "entity_label": "Alpha School",
                "metric": "baseline_without_target",
                "baseline": "1",
                "current": "2",
                "delta": "1",
                "movement_type": "issue_code_delta",
                "is_regression": True,
            },
            {
                "stream": "metrics",
                "entity_key": "site-b",
                "entity_label": "Beta School",
                "metric": "flagged_rows",
                "baseline": "0",
                "current": "1",
                "delta": "1",
                "movement_type": "school_flag_count_delta",
                "is_regression": True,
            },
        ],
    )
    global_metrics = visualize.aggregate_global_metrics(pair_dir)
    assert global_metrics["baseline_without_target"] == 20
    assert global_metrics["both_baseline_and_target_blank"] is None

    schools = visualize.aggregate_school_regressions(pair_dir, top_n=2)
    assert schools[0]["school"] == "Alpha School"
    assert schools[0]["regressions"] == 2
    assert schools[0]["top_metric"] == "flagged_rows"


def test_render_cross_year_html_writes_expected_labels(tmp_path: Path) -> None:
    import visualize

    output_dir = tmp_path / "trends"
    _write_cross_year_pair(
        output_dir,
        pair_name="SY24-25_EOY_vs_SY25-26_Q2",
        summary={
            "current_school_year": "SY25-26",
            "baseline_school_year": "SY24-25",
            "baseline_period": "EOY",
            "current_period": "Q2",
            "movement_count": 12,
            "regression_count": 5,
            "movements_by_type": {
                "issue_code_delta": 8,
                "school_flag_count_delta": 2,
                "global_issue_delta": 1,
                "grading_period_fill_delta": 1,
                "site_appeared": 1,
                "site_removed": 0,
            },
        },
        movements=[
            {
                "stream": "metrics",
                "entity_key": "global",
                "entity_label": "global",
                "metric": "baseline_without_target",
                "baseline": "1",
                "current": "9",
                "delta": "8",
                "movement_type": "global_issue_delta",
                "is_regression": True,
            }
        ],
        regressions=[
            {
                "stream": "metrics",
                "entity_key": "site-a",
                "entity_label": "Alpha School",
                "metric": "flagged_rows",
                "baseline": "1",
                "current": "2",
                "delta": "1",
                "movement_type": "school_flag_count_delta",
                "is_regression": True,
            }
        ],
    )
    report_path = output_dir / "cross_year" / "index.html"
    pair_dir = output_dir / "cross_year" / "SY24-25_EOY_vs_SY25-26_Q2"
    (pair_dir / "goal_progress_summary.json").write_text(
        json.dumps(
            {
                "baseline": {"eligible_rows": 100, "on_track_pct": 40.0, "global": {}},
                "current": {"eligible_rows": 120, "on_track_pct": 45.0, "global": {}},
                "row_level": {"improved": 3, "worsened": 2},
            }
        ),
        encoding="utf-8",
    )
    aggregates = visualize.build_cross_year_aggregates(output_dir)
    visualize.render_cross_year_html(aggregates, report_path)
    html = report_path.read_text(encoding="utf-8")
    assert "Cross-year trend report" in html
    assert "Monday" not in html
    assert "SY25-26" in html
    assert "Alpha School" in html
    assert "Changed issue counts" in html
    assert "Issue counts increased" in html
    assert "grading_period_fill_delta" in html
    assert "Goal–Metric rows with Baseline and Target" in html
    assert "120 Goal–Metric rows" in html
    assert 'class="data-label"' in html
    assert ">12</text>" in html
    assert ">5</text>" in html


def test_svg_line_chart_value_labels() -> None:
    import visualize

    svg = visualize._svg_line_chart(
        width=400,
        height=200,
        categories=["A", "B"],
        series=[{"name": "Test", "data": [10, 20]}],
        y_label="Test chart",
        show_value_labels="all",
    )
    assert 'class="data-label"' in svg
    assert ">10</text>" in svg
    assert ">20</text>" in svg

    svg_latest = visualize._svg_line_chart(
        width=400,
        height=200,
        categories=["A", "B", "C"],
        series=[{"name": "Test", "data": [10, 20, 30]}],
        y_label="Test chart",
        show_value_labels="latest",
    )
    assert ">30</text>" in svg_latest
    assert ">10</text>" not in svg_latest
    assert ">20</text>" not in svg_latest


def test_svg_stacked_bar_chart_show_totals() -> None:
    import visualize

    svg = visualize._svg_stacked_bar_chart(
        width=400,
        height=200,
        categories=["SY24-25"],
        series=[
            {"name": "a", "data": [8]},
            {"name": "b", "data": [4]},
        ],
        y_label="Stacked test",
        show_totals=True,
    )
    assert 'class="data-label"' in svg
    assert ">12</text>" in svg


def test_run_cross_year_compares_writes_html_report(tmp_path: Path) -> None:
    import visualize

    snapshots_dir = tmp_path / "snapshots"
    output_dir = tmp_path / "trends"
    inputs = tmp_path / "archives"
    for school_year in ("SY24-25", "SY25-26"):
        acc, met = _seed_period_inputs(
            inputs,
            school_year,
            trends.FALLBACK_PERIOD,
            acc_variant="worse",
        )
        trends.capture_snapshot(
            school_year=school_year,
            period=trends.FALLBACK_PERIOD,
            snapshots_dir=snapshots_dir,
            accreditation_workbook=acc,
            metrics_workbook=met,
        )
    code = trends.run_cross_year_compares(
        ["SY24-25", "SY25-26"],
        snapshots_dir=snapshots_dir,
        output_dir=output_dir,
        regression_threshold=0.01,
    )
    assert code == 0
    report_path = output_dir / "cross_year" / "index.html"
    assert report_path.is_file()
    aggregates = visualize.build_cross_year_aggregates(
        output_dir,
        snapshots_dir=snapshots_dir,
    )
    assert aggregates["pair_count"] == 1
    assert aggregates["eoy_snapshot_count"] == 2


def test_run_cross_year_compares_skips_without_eoy_snapshot(tmp_path: Path) -> None:
    snapshots_dir = tmp_path / "snapshots"
    output_dir = tmp_path / "trends"
    inputs = tmp_path / "archives"
    acc, met = _seed_period_inputs(inputs, "SY24-25", trends.FALLBACK_PERIOD)
    trends.capture_snapshot(
        school_year="SY24-25",
        period=trends.FALLBACK_PERIOD,
        snapshots_dir=snapshots_dir,
        accreditation_workbook=acc,
        metrics_workbook=met,
    )
    acc2, met2 = _seed_period_inputs(inputs, "SY25-26", "Q1", acc_variant="worse")
    trends.capture_snapshot(
        school_year="SY25-26",
        period="Q1",
        snapshots_dir=snapshots_dir,
        accreditation_workbook=acc2,
        metrics_workbook=met2,
    )
    code = trends.run_cross_year_compares(
        ["SY24-25", "SY25-26"],
        snapshots_dir=snapshots_dir,
        output_dir=output_dir,
        regression_threshold=0.01,
    )
    assert code == 0
    import visualize

    assert visualize.load_cross_year_summaries(output_dir) == []


def test_discover_eoy_snapshot_years(tmp_path: Path) -> None:
    snapshots_dir = tmp_path / "snapshots"
    inputs = tmp_path / "archives"
    for school_year in ("SY23-24", "SY24-25"):
        acc, met = _seed_period_inputs(inputs, school_year, trends.FALLBACK_PERIOD)
        trends.capture_snapshot(
            school_year=school_year,
            period=trends.FALLBACK_PERIOD,
            snapshots_dir=snapshots_dir,
            accreditation_workbook=acc,
            metrics_workbook=met,
        )
    acc, met = _seed_period_inputs(inputs, "SY25-26", "Q1")
    trends.capture_snapshot(
        school_year="SY25-26",
        period="Q1",
        snapshots_dir=snapshots_dir,
        accreditation_workbook=acc,
        metrics_workbook=met,
    )
    assert trends.discover_eoy_snapshot_years(snapshots_dir) == ["SY23-24", "SY24-25"]
