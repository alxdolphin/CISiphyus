from __future__ import annotations

import hashlib
import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def sha256_12(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()[:12]


def normalize_cell(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


_YEAR_NUMBER_RE = re.compile(r"\d{2,4}")


def school_year_label(value: Any) -> str | None:
    # WHY: exports write the year as 2026-2027, 2026/27, or SY26-27; compare on SYxx-yy
    numbers = _YEAR_NUMBER_RE.findall(normalize_cell(value))
    if len(numbers) < 2:
        return None
    return f"SY{numbers[0][-2:]}-{numbers[1][-2:]}"


def check_school_year_column(
    workbook_path: Path,
    sheet_name: str,
    header_row: int,
    column: str,
    expected: str,
) -> tuple[bool, str, list[str]]:
    expected_label = school_year_label(expected) or expected
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        sheet = workbook[sheet_name]
        header = next(
            sheet.iter_rows(min_row=header_row, max_row=header_row, values_only=True), ()
        )
        column_index = [normalize_cell(cell) for cell in header].index(column.strip())
        seen: set[str] = set()
        offending: list[str] = []
        for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
            raw = normalize_cell(row[column_index]) if column_index < len(row) else ""
            if not raw or raw in seen:
                continue
            seen.add(raw)
            if school_year_label(raw) != expected_label:
                offending.append(raw)
    finally:
        workbook.close()
    if offending:
        return False, "school_year_mismatch", sorted(offending)
    if not seen:
        # WHY: a fresh school year can export zero rows; nothing to contradict
        return True, "school_year_check_skipped_no_rows", []
    return True, "school_year_check_passed", []


def find_header_row(
    workbook_path: Path,
    required_columns: list[str],
    max_scan_rows: int = 30,
) -> dict[str, Any]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    required = {column.strip() for column in required_columns}

    try:
        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            for row_index, row in enumerate(
                sheet.iter_rows(min_row=1, max_row=max_scan_rows, values_only=True),
                start=1,
            ):
                values = [normalize_cell(cell) for cell in row]
                present = {value for value in values if value}
                missing = sorted(required - present)
                if not missing:
                    return {
                        "workbook_opened": True,
                        "sheet_name": sheet_name,
                        "header_row": row_index,
                        "required_columns_present": True,
                        "missing_required_columns": [],
                        "detected_columns": values,
                    }

        first_sheet = workbook[workbook.sheetnames[0]]
        first_row = next(first_sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        return {
            "workbook_opened": True,
            "sheet_name": None,
            "header_row": None,
            "required_columns_present": False,
            "missing_required_columns": sorted(required),
            "detected_columns_first_row": [normalize_cell(cell) for cell in first_row],
        }
    finally:
        workbook.close()


def validate_workbook(
    workbook_path: Path,
    required_columns: list[str],
    min_size_bytes: int,
    school_year_check: tuple[str, str] | None = None,
) -> tuple[bool, str, dict[str, Any]]:
    metadata: dict[str, Any] = {
        "path": str(workbook_path),
        "size_bytes": workbook_path.stat().st_size,
        "min_size_bytes": min_size_bytes,
        "sha256_12": sha256_12(workbook_path),
    }

    if workbook_path.stat().st_size < min_size_bytes:
        return False, "artifact_too_small", metadata

    try:
        header_columns = list(required_columns)
        if school_year_check and school_year_check[0] not in header_columns:
            header_columns.append(school_year_check[0])
        header_result = find_header_row(workbook_path, header_columns)
        metadata.update(header_result)
        if not header_result.get("required_columns_present"):
            return False, "required_columns_missing", metadata
        if school_year_check:
            column, expected = school_year_check
            ok, reason, offending = check_school_year_column(
                workbook_path,
                header_result["sheet_name"],
                header_result["header_row"],
                column,
                expected,
            )
            metadata["school_year_check"] = {
                "column": column,
                "expected": expected,
                "reason": reason,
                "offending_values": offending,
            }
            if not ok:
                return False, reason, metadata
        return True, "validation_passed", metadata
    except Exception as exc:
        metadata["exception_type"] = type(exc).__name__
        metadata["exception_message"] = str(exc)
        return False, "workbook_open_failed", metadata
