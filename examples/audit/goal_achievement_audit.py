#!/usr/bin/env python3
"""Goal achievement audit: GAR on goal tracking + drilldown cross-check."""

from __future__ import annotations

import argparse
import csv
import json
import os
import sys
from collections import Counter
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

_EXAMPLES_DIR = Path(__file__).resolve().parent.parent
for _subdir in ("audit", "goal_achievement", "accreditation"):
    _path = _EXAMPLES_DIR / _subdir
    if str(_path) not in sys.path:
        sys.path.insert(0, str(_path))

import audit


def default_output_dir() -> Path:
    explicit = (os.environ.get("GOAL_ACHIEVEMENT_OUTPUT_DIR") or "").strip()
    if explicit:
        return Path(explicit).expanduser().resolve()
    return Path(__file__).resolve().parents[2] / "artifacts" / "goal_achievement"


def default_local_inputs_dir() -> Path:
    explicit = (os.environ.get("GOAL_ACHIEVEMENT_LOCAL_INPUTS_DIR") or "").strip()
    if explicit:
        return Path(explicit).expanduser().resolve()
    return default_output_dir()


def goal_achievement_audit_archive_dir() -> Path:
    return Path(__file__).resolve().parents[2] / "artifacts" / "goal_achievement_audit"


def preferred_goal_progress_workbook(*, school_year: str | None = None) -> Path:
    explicit = (os.environ.get("GOAL_ACHIEVEMENT_GOAL_PROGRESS_WORKBOOK") or "").strip()
    if explicit:
        return Path(explicit).expanduser().resolve()
    if school_year:
        resolved = audit.resolve_goal_progress_workbook(school_year)
        if resolved is not None:
            return resolved
        named = default_local_inputs_dir() / f"{school_year}_GoalProgress.xlsx"
        return named
    return default_local_inputs_dir() / "SY25-26_GoalProgress.xlsx"


def preferred_student_metrics_workbook(*, school_year: str | None = None) -> Path:
    explicit = (os.environ.get("GOAL_ACHIEVEMENT_STUDENT_METRICS_WORKBOOK") or "").strip()
    if explicit:
        return Path(explicit).expanduser().resolve()
    base = default_local_inputs_dir()
    if school_year:
        named = base / f"{school_year}_StudentMetricsSummary.xlsx"
        if named.is_file():
            return named
    return base / "SY25-26_StudentMetricsSummary.xlsx"


def preferred_accreditation_workbook() -> Path:
    explicit = (os.environ.get("GOAL_ACHIEVEMENT_ACCREDITATION_WORKBOOK") or "").strip()
    if explicit:
        return Path(explicit).expanduser().resolve()
    return default_local_inputs_dir() / "Accreditation_Report.xlsx"


@dataclass(frozen=True)
class AuditInputs:
    goal_progress_workbook: Path
    student_metrics_workbook: Path
    accreditation_workbook: Path


def load_drilldown_rows(workbook_path: Path) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if audit.DRILLDOWN_SHEET_NAME not in workbook.sheetnames:
            return []
        worksheet = workbook[audit.DRILLDOWN_SHEET_NAME]
        header_row: list[str] | None = None
        header_index = -1
        for row_number, row in enumerate(
            worksheet.iter_rows(values_only=True, max_row=10),
            start=1,
        ):
            normalized = [audit._normalize_header_cell(value) for value in row]
            if "Student ID" in normalized:
                header_row = normalized
                header_index = row_number
                break
        if header_row is None:
            return []
        records: list[dict[str, Any]] = []
        for excel_row, row_values in enumerate(
            worksheet.iter_rows(min_row=header_index + 1, values_only=True),
            start=header_index + 1,
        ):
            if not row_values or not any(row_values):
                continue
            record: dict[str, Any] = {"drilldown_excel_row": excel_row}
            for index, header in enumerate(header_row):
                if not header:
                    continue
                value = row_values[index] if index < len(row_values) else None
                record[header] = audit.clean_value(str(value) if value is not None else None)
            records.append(record)
        return records
    finally:
        workbook.close()


def _intish(value: object) -> int:
    text = audit.clean_value(str(value) if value is not None else None) or "0"
    try:
        return int(float(text))
    except ValueError:
        return 0


def _metrics_by_student(metrics_rows: list[dict[str, Any]]) -> dict[str, list[dict[str, Any]]]:
    grouped: dict[str, list[dict[str, Any]]] = {}
    for row in metrics_rows:
        student_id = audit.clean_value(str(row.get("Student ID") or "")) or ""
        if not student_id:
            continue
        grouped.setdefault(student_id, []).append(row)
    return grouped


def _has_latest_progress(row: dict[str, Any]) -> bool:
    return bool(audit.clean_value(str(row.get("Latest Progress") or "")))


def _has_target(row: dict[str, Any]) -> bool:
    return bool(audit.clean_value(str(row.get("Target") or "")))


def evaluate_cross_audit(
    drilldown_rows: list[dict[str, Any]],
    metrics_rows: list[dict[str, Any]],
) -> dict[str, Any]:
    metrics_by_student = _metrics_by_student(metrics_rows)
    drilldown_students = {
        audit.clean_value(str(row.get("Student ID") or "")) or ""
        for row in drilldown_rows
    }
    exception_rows: list[dict[str, Any]] = []
    summary: Counter[str] = Counter()

    for row in drilldown_rows:
        student_id = audit.clean_value(str(row.get("Student ID") or "")) or ""
        if not student_id:
            continue
        school = audit.clean_value(str(row.get("School") or "")) or ""
        student_metrics = metrics_by_student.get(student_id, [])
        metric_rows = len(student_metrics)
        metrics_with_latest_progress = sum(1 for item in student_metrics if _has_latest_progress(item))
        metrics_with_target = sum(1 for item in student_metrics if _has_target(item))
        achievement_count = _intish(row.get("# of Metrics with Goal Achievement entered"))
        all_entered = (audit.clean_value(str(row.get(
            "Goal Achievement Entered for ALL assigned goals?"
        ) or "")) or "").lower()
        assigned = _intish(row.get("# ABCS Metrics Assigned")) + _intish(
            row.get("# Non-ABCS Metrics Assigned")
        )
        base = {
            "audit_layer": "drilldown",
            "student_id": student_id,
            "school": school,
            "assigned_metrics_drilldown": assigned,
            "achievement_count_drilldown": achievement_count,
            "all_achievement_entered_drilldown": row.get(
                "Goal Achievement Entered for ALL assigned goals?"
            ),
            "metric_rows_student_metrics": metric_rows,
            "metrics_with_latest_progress": metrics_with_latest_progress,
            "metrics_with_target": metrics_with_target,
            "drilldown_excel_row": row.get("drilldown_excel_row"),
        }

        if achievement_count != metrics_with_latest_progress:
            exception_rows.append(
                {
                    **base,
                    "exception_type": "achievement_count_mismatch",
                    "reason": (
                        f"drilldown achievement count {achievement_count} != "
                        f"metrics rows with Latest Progress {metrics_with_latest_progress}"
                    ),
                }
            )
            summary["achievement_count_mismatch"] += 1

        if all_entered == "yes" and metrics_with_latest_progress < metric_rows:
            exception_rows.append(
                {
                    **base,
                    "exception_type": "all_entered_missing_progress",
                    "reason": (
                        "drilldown marks all achievement entered but some metrics "
                        "lack Latest Progress"
                    ),
                }
            )
            summary["all_entered_missing_progress"] += 1

        if all_entered != "yes" and metrics_with_latest_progress > 0:
            exception_rows.append(
                {
                    **base,
                    "exception_type": "all_progress_not_marked_complete",
                    "reason": (
                        "drilldown does not mark all achievement entered but metrics "
                        "rows have Latest Progress"
                    ),
                }
            )
            summary["all_progress_not_marked_complete"] += 1

    metrics_only = set(metrics_by_student) - {sid for sid in drilldown_students if sid}
    summary["drilldown_students"] = len(drilldown_rows)
    summary["students_with_metrics_rows"] = len(metrics_by_student)
    summary["metrics_only_students"] = len(metrics_only)

    return {
        "exception_rows": exception_rows,
        "summary": dict(summary),
        "drilldown_row_count": len(drilldown_rows),
        "metrics_row_count": len(metrics_rows),
    }


def _filter_gar_records(
    records: list[dict[str, Any]],
    *,
    school: str | None,
    metric: str | None,
    excluded_student_ids: set[str],
) -> list[dict[str, Any]]:
    filtered: list[dict[str, Any]] = []
    for record in records:
        student_id = audit.clean_value(str(record.get("Student ID") or "")) or ""
        if student_id and student_id in excluded_student_ids:
            continue
        if school and audit.clean_value(str(record.get("Home School") or "")) != school:
            continue
        if metric and audit.clean_value(str(record.get("Metric") or "")) != metric:
            continue
        filtered.append(record)
    return filtered


def run_inherited_gar_audit(
    records: list[dict[str, Any]],
    *,
    school: str | None,
    metric: str | None,
) -> dict[str, Any]:
    """Backward-compatible alias for tests."""
    return audit.run_gar_audit(records, school=school, metric=metric)


def run_inherited_metrics_audit(
    metrics_module: Any,
    metrics_rows: list[dict[str, Any]],
    *,
    include_domain_scale: bool,
    ramp_days_after_enrollment: int,
    quarter_end_grace_days: int,
    as_of_date: Any,
) -> dict[str, Any]:
    return metrics_module.run_audit(
        metrics_rows,
        include_domain_scale=include_domain_scale,
        ramp_days_after_enrollment=ramp_days_after_enrollment,
        quarter_end_grace_days=quarter_end_grace_days,
        as_of_date=as_of_date,
    )


def _write_csv(path: Path, rows: list[dict[str, Any]], fieldnames: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def _render_summary_markdown(payload: dict[str, Any]) -> str:
    gar_summary = (payload.get("gar_audit") or {}).get("summary") or {}
    lines = [
        "# Goal Achievement Audit",
        "",
        "GAR on goal tracking (`CIS_StudentProgress_Detail`) plus accreditation drilldown cross-check.",
        "",
        f"- Goal progress workbook: `{payload.get('goal_progress_workbook')}`",
        f"- Student metrics workbook: `{payload.get('student_metrics_workbook')}`",
        f"- Accreditation workbook: `{payload.get('accreditation_workbook')}`",
        f"- GAR rows analyzed: {payload.get('gar_row_count', 0)}",
        f"- Metrics rows (drilldown): {payload.get('metrics_row_count', 0)}",
        f"- Drilldown rows: {payload.get('drilldown_row_count', 0)}",
        "",
        "## Exception totals",
    ]
    totals = payload.get("exception_totals") or {}
    lines.append(f"- GAR exceptions: {totals.get('gar', 0)}")
    lines.append(f"- Drilldown cross-check: {totals.get('drilldown', 0)}")
    lines.extend(["", "## GAR audit"])
    lines.append(f"- Matched: {gar_summary.get('matched', 0)}")
    lines.append(f"- Mismatched: {gar_summary.get('mismatched', 0)}")
    lines.append(f"- Manual review rows: {gar_summary.get('manual_review_rows', 0)}")
    for key, value in sorted(gar_summary.items()):
        if key not in {"matched", "mismatched", "manual_review_rows", "total_rows"}:
            lines.append(f"- `{key}`: {value}")
    lines.extend(["", "## Drilldown cross-check"])
    drilldown_summary = (payload.get("drilldown_cross_audit") or {}).get("summary") or {}
    for key, value in sorted(drilldown_summary.items()):
        lines.append(f"- `{key}`: {value}")
    return "\n".join(lines) + "\n"


def resolve_inputs(
    *,
    school_year: str | None = None,
    force_fetch: bool = False,
) -> AuditInputs:
    import accreditation

    goal_progress_path = preferred_goal_progress_workbook(school_year=school_year)
    metrics_path = preferred_student_metrics_workbook(school_year=school_year)
    accred_path = preferred_accreditation_workbook()
    audit.fetch_goal_progress_workbook(
        destination=goal_progress_path,
        force_fetch=force_fetch,
        school_year=school_year,
    )
    audit.fetch_student_metrics_workbook(
        destination=metrics_path,
        force_fetch=force_fetch,
        school_year=school_year,
    )
    accreditation.fetch_accreditation_workbook(
        destination=accred_path,
        force_fetch=force_fetch,
        school_year=school_year,
    )
    return AuditInputs(
        goal_progress_workbook=goal_progress_path,
        student_metrics_workbook=metrics_path,
        accreditation_workbook=accred_path,
    )


def run_goal_achievement_audit(
    inputs: AuditInputs,
    *,
    school_year: str | None = None,
    school: str | None = None,
    metric: str | None = None,
) -> dict[str, Any]:
    gar_records = audit.load_goal_progress_gar_records(inputs.goal_progress_workbook)
    excluded = audit.resolve_ewgspe_only_student_ids(
        school_year,
        goal_progress_workbook=inputs.goal_progress_workbook,
    )
    gar_records = _filter_gar_records(
        gar_records,
        school=school,
        metric=metric,
        excluded_student_ids=excluded,
    )
    gar_results = audit.run_gar_audit(
        gar_records,
        school=school,
        metric=metric,
    )

    metrics_rows = audit.load_student_metrics_workbook(inputs.student_metrics_workbook)
    drilldown_rows = load_drilldown_rows(inputs.accreditation_workbook)
    drilldown_audit = evaluate_cross_audit(drilldown_rows, metrics_rows)

    gar_payload = {
        "workbook": str(inputs.goal_progress_workbook.resolve()),
        "sheet_name": gar_results.get("sheet_name", audit.GOAL_PROGRESS_SHEET),
        "school_filter": school,
        "metric_filter": metric,
        "record_count": gar_results.get("record_count", len(gar_records)),
        "report_details": gar_results.get("report_details") or {},
        "summary": gar_results.get("summary") or {},
        "direction_counts": gar_results.get("direction_counts") or {},
        "mismatch_examples": gar_results.get("mismatch_examples") or [],
    }

    return {
        "run_at": datetime.now(timezone.utc).isoformat(),
        "school_year": school_year,
        "goal_progress_workbook": str(inputs.goal_progress_workbook.resolve()),
        "student_metrics_workbook": str(inputs.student_metrics_workbook.resolve()),
        "accreditation_workbook": str(inputs.accreditation_workbook.resolve()),
        "gar_row_count": gar_results.get("record_count", len(gar_records)),
        "metrics_row_count": len(metrics_rows),
        "drilldown_row_count": len(drilldown_rows),
        "gar_audit": {
            **gar_payload,
            "exception_rows": gar_results.get("exception_rows") or [],
        },
        "drilldown_cross_audit": drilldown_audit,
        "exception_totals": {
            "gar": len(gar_results.get("exception_rows") or []),
            "drilldown": len(drilldown_audit.get("exception_rows") or []),
        },
    }


def export_goal_achievement_results(payload: dict[str, Any], output_dir: Path) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    gar_audit = payload.get("gar_audit") or {}
    gar_rows = gar_audit.get("exception_rows") or []
    drilldown_rows = (payload.get("drilldown_cross_audit") or {}).get("exception_rows") or []

    audit.gar_write_csv_report(gar_rows, output_dir / "gar_exceptions.csv")
    audit.gar_write_markdown_report(gar_audit, output_dir / "gar_audit_report.md")
    _write_csv(
        output_dir / "drilldown_exceptions.csv",
        drilldown_rows,
        [
            "audit_layer",
            "exception_type",
            "student_id",
            "school",
            "assigned_metrics_drilldown",
            "achievement_count_drilldown",
            "all_achievement_entered_drilldown",
            "metric_rows_student_metrics",
            "metrics_with_latest_progress",
            "metrics_with_target",
            "reason",
            "drilldown_excel_row",
        ],
    )

    school_year = payload.get("school_year")
    if school_year:
        archive_dir = goal_achievement_audit_archive_dir()
        archive_dir.mkdir(parents=True, exist_ok=True)
        audit.gar_write_csv_report(
            gar_rows,
            archive_dir / f"{school_year}_GoalAchievement_AUDIT.csv",
        )

    (output_dir / "goal_achievement_summary.json").write_text(
        json.dumps(payload, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    (output_dir / "goal_achievement_summary.md").write_text(
        _render_summary_markdown(payload),
        encoding="utf-8",
    )


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description=(
            "Goal achievement audit: GAR on goal_tracking_student_goals export, "
            "plus accreditation drilldown cross-check."
        ),
    )
    parser.add_argument("--school-year", metavar="SY", help="School year label (e.g. SY25-26).")
    parser.add_argument(
        "--goal-progress-workbook",
        type=Path,
        help="Path to goal tracking student goals export.",
    )
    parser.add_argument(
        "--student-metrics-workbook",
        type=Path,
        help="Path to student metrics summary export (drilldown cross-check).",
    )
    parser.add_argument(
        "--accreditation-workbook",
        type=Path,
        help="Path to accreditation report export.",
    )
    parser.add_argument("--school", help="Optional exact Home School filter for GAR.")
    parser.add_argument("--metric", help="Optional exact Metric filter for GAR.")
    parser.add_argument(
        "--force-fetch",
        action="store_true",
        help="Re-export inputs from CISDM even when local copies are fresh.",
    )
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    school_year = audit.normalize_school_year(args.school_year) if args.school_year else None
    if (
        args.goal_progress_workbook
        and args.student_metrics_workbook
        and args.accreditation_workbook
    ):
        inputs = AuditInputs(
            goal_progress_workbook=args.goal_progress_workbook.expanduser().resolve(),
            student_metrics_workbook=args.student_metrics_workbook.expanduser().resolve(),
            accreditation_workbook=args.accreditation_workbook.expanduser().resolve(),
        )
    else:
        inputs = resolve_inputs(school_year=school_year, force_fetch=args.force_fetch)

    payload = run_goal_achievement_audit(
        inputs,
        school_year=school_year,
        school=args.school,
        metric=args.metric,
    )
    export_goal_achievement_results(payload, default_output_dir())
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
