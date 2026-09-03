#!/usr/bin/env python3
"""
Student metrics audit for CISDM student metrics summary exports.

Run (fetches from CISDM when needed):
  cisiphyus audit metrics
  cisiphyus audit metrics --year SY24-25
  cisiphyus audit metrics --year all

Pin a local workbook:
  cisiphyus audit metrics --workbook path/to/StudentMetricsSummary.xlsx

Per-year outputs: artifacts/audit/<school-year>/
`--year all` writes a cumulative comparison to artifacts/audit/.
"""

from __future__ import annotations

import argparse
import csv
import json
import os
import re
import shutil
import subprocess
import sys
from collections import Counter, defaultdict
from dataclasses import asdict, dataclass
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any, Callable, Literal, TypeAlias

from openpyxl import load_workbook

SCRIPT_DIR = Path(__file__).resolve().parent


def default_output_dir() -> Path:
    explicit = (os.environ.get("AUDIT_OUTPUT_DIR") or "").strip()
    if explicit:
        return Path(explicit).expanduser().resolve()
    return Path(__file__).resolve().parents[2] / "artifacts" / "audit"

# --- cisiphyus fetch (live export when --workbook is omitted) ---

TRUE_VALUES = frozenset({"1", "true", "yes", "on"})
STUDENT_METRICS_MAX_AGE_HOURS = 24.0
CISIPHYUS_REPORT_ID = "student_metrics_summary"
GOAL_TRACKING_REPORT_ID = "goal_tracking_student_goals"
DEFAULT_STUDENT_METRICS_FILENAME = "SY25-26_StudentMetricsSummary.xlsx"
SCHOOL_YEAR_LABEL_RE = re.compile(r"^SY\d{2}-\d{2}$", re.IGNORECASE)
SCHOOL_YEAR_SHORT_RE = re.compile(r"^SY\d{2}$", re.IGNORECASE)
ARCHIVE_PERIOD = "EOY"
METRICS_GLOBS = (
    "metrics.xlsx",
    "*StudentMetrics*.xlsx",
    "*student_metrics*.xlsx",
    "*MetricsSummary*.xlsx",
)
CUMULATIVE_METRIC_ORDER = (
    "rows_evaluated",
    "baseline_and_target_present",
    "baseline_without_target_excl_supplemental",
    "partial_goal_metrics_flagged",
    "accepted_domain_tracking",
    "accepted_supplemental",
    "accepted_no_abc_goals",
    "accepted_exceptions_total",
    "rows_flagged",
    "direction_mismatches",
    "scale_mismatches",
)


@dataclass(frozen=True)
class FetchResult:
    destination: Path
    triggered: bool
    succeeded: bool
    reason: str


def _flag_true(raw: str | None) -> bool:
    return (raw or "").strip().lower() in TRUE_VALUES


def _default_cisiphyus_root() -> Path:
    # WHY: prototypical app lives at examples/audit/ inside the cisiphyus repo
    return Path(__file__).resolve().parents[2]


def _default_local_inputs_dir() -> Path:
    return Path(__file__).resolve().parents[2] / "artifacts" / "audit"


def default_archives_dir() -> Path:
    explicit = (os.environ.get("AUDIT_ARCHIVES_DIR") or os.environ.get("TREND_INPUTS_DIR") or "").strip()
    if explicit:
        return Path(explicit).expanduser().resolve()
    return Path(__file__).resolve().parents[2] / "artifacts" / "archives"


def _load_config_programs() -> dict[str, int]:
    src_dir = _default_cisiphyus_root() / "src"
    if str(src_dir) not in sys.path:
        sys.path.insert(0, str(src_dir))
    import config

    return config.load_school_year_programs()


def current_school_year() -> str | None:
    try:
        programs = _load_config_programs()
    except (FileNotFoundError, ValueError, OSError):
        return None
    import config

    return config.default_school_year(programs)


def student_metrics_workbook_name(school_year: str) -> str:
    return f"{school_year}_StudentMetricsSummary.xlsx"


def default_student_metrics_filename() -> str:
    school_year = current_school_year()
    if school_year:
        return student_metrics_workbook_name(school_year)
    return DEFAULT_STUDENT_METRICS_FILENAME


def discover_audit_school_years(*, include_config_years: bool = True) -> list[str]:
    years: set[str] = set()
    archives_dir = default_archives_dir()
    if archives_dir.is_dir():
        for child in archives_dir.iterdir():
            if child.is_dir() and child.name.upper().startswith("SY"):
                years.add(child.name)
    if include_config_years:
        try:
            years.update(_load_config_programs().keys())
        except (FileNotFoundError, ValueError, OSError):
            pass
    import config

    return sorted(years, key=config._school_year_sort_key)


def normalize_school_year(token: str, known_years: list[str] | None = None) -> str:
    stripped = token.strip()
    if not stripped:
        raise ValueError("school year cannot be empty")
    upper = stripped.upper()
    catalog = known_years if known_years is not None else discover_audit_school_years()
    by_upper = {year.upper(): year for year in catalog}
    if SCHOOL_YEAR_LABEL_RE.match(upper):
        if upper in by_upper:
            return by_upper[upper]
        return stripped
    if SCHOOL_YEAR_SHORT_RE.match(upper):
        matches = [year for year in catalog if year.upper().startswith(f"{upper}-")]
        if len(matches) == 1:
            return matches[0]
        if len(matches) > 1:
            joined = ", ".join(matches)
            raise ValueError(f"ambiguous school year {token!r}: {joined}")
        raise ValueError(f"unknown school year {token!r}")
    raise ValueError(f"invalid school year {token!r}; expected SY24-25 or SY24")


def _find_metrics_workbook(period_dir: Path) -> Path | None:
    for pattern in METRICS_GLOBS:
        matches = sorted(period_dir.glob(pattern))
        if matches:
            return matches[0].resolve()
    return None


ACCREDITATION_REPORT_FILENAME = "Accreditation_Report.xlsx"
DRILLDOWN_SHEET_NAME = "Accreditation Student Drilldown"
FINAL_GOAL_ACHIEVEMENT_COLUMN = (
    "Goal Achievement Entered for ALL assigned goals?"
)
STUDENT_ID_COLUMN = "Student ID"


def _normalize_header_cell(value: object) -> str:
    return " ".join(str(value or "").split())


def resolve_accreditation_workbook_for_school_year(school_year: str | None) -> Path | None:
    repo_root = Path(__file__).resolve().parents[2]
    live_accreditation = (
        repo_root / "artifacts" / "accreditation" / ACCREDITATION_REPORT_FILENAME
    )
    pull_accreditation = (
        default_archives_dir()
        / school_year
        / "pulls"
        / "accreditation"
        / "raw.xlsx"
        if school_year
        else None
    )
    archive_accreditation = (
        default_archives_dir()
        / school_year
        / ARCHIVE_PERIOD
        / ACCREDITATION_REPORT_FILENAME
        if school_year
        else None
    )
    current_school_year_label = current_school_year()
    candidates: list[Path] = []
    if school_year and current_school_year_label and school_year == current_school_year_label:
        if live_accreditation.is_file():
            candidates.append(live_accreditation)
        if pull_accreditation is not None and pull_accreditation.is_file():
            candidates.append(pull_accreditation)
        if archive_accreditation is not None and archive_accreditation.is_file():
            candidates.append(archive_accreditation)
    elif school_year:
        if pull_accreditation is not None and pull_accreditation.is_file():
            candidates.append(pull_accreditation)
        if archive_accreditation is not None and archive_accreditation.is_file():
            candidates.append(archive_accreditation)
    elif live_accreditation.is_file():
        candidates.append(live_accreditation)
    for path in candidates:
        return path.resolve()
    return None


def load_final_goal_achievement_student_ids(workbook_path: Path) -> set[str]:
    from openpyxl import load_workbook

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if DRILLDOWN_SHEET_NAME not in workbook.sheetnames:
            return set()
        worksheet = workbook[DRILLDOWN_SHEET_NAME]
        header_row: list[str] | None = None
        header_index = -1
        for row_number, row in enumerate(
            worksheet.iter_rows(values_only=True, max_row=10),
            start=1,
        ):
            normalized = [_normalize_header_cell(value) for value in row]
            if STUDENT_ID_COLUMN in normalized and FINAL_GOAL_ACHIEVEMENT_COLUMN in normalized:
                header_row = normalized
                header_index = row_number
                break
        if header_row is None:
            return set()
        student_index = header_row.index(STUDENT_ID_COLUMN)
        achievement_index = header_row.index(FINAL_GOAL_ACHIEVEMENT_COLUMN)
        completed: set[str] = set()
        for row in worksheet.iter_rows(min_row=header_index + 1, values_only=True):
            if not row or not any(row):
                continue
            if student_index >= len(row) or achievement_index >= len(row):
                continue
            student_id = clean_value(
                str(row[student_index]) if row[student_index] is not None else None
            )
            achievement = clean_value(
                str(row[achievement_index]) if row[achievement_index] is not None else None
            )
            if student_id and achievement and achievement.lower() == "yes":
                completed.add(student_id)
        return completed
    finally:
        workbook.close()


def resolve_final_goal_achievement_students(school_year: str | None) -> set[str] | None:
    workbook = resolve_accreditation_workbook_for_school_year(school_year)
    if workbook is None:
        return None
    return load_final_goal_achievement_student_ids(workbook)


def default_goal_achievement_datasets_dir() -> Path:
    explicit = (os.environ.get("GOAL_ACHIEVEMENT_DATASETS_DIR") or "").strip()
    if explicit:
        return Path(explicit).expanduser().resolve()
    cis_root = _default_cisiphyus_root().parent.parent
    return cis_root / "data" / "datasets" / "Goal Achievement"


def resolve_goal_progress_workbook(school_year: str | None) -> Path | None:
    if not school_year:
        return None
    archives_dir = default_archives_dir()
    repo_root = _default_cisiphyus_root()
    candidates = [
        archives_dir
        / school_year
        / "pulls"
        / "goal_progress"
        / "raw.xlsx",
        archives_dir
        / school_year
        / "pulls"
        / "goal_tracking_student_goals"
        / "raw.xlsx",
        default_goal_achievement_datasets_dir()
        / GOAL_PROGRESS_FILENAME_TEMPLATE.format(school_year=school_year),
    ]
    current = current_school_year()
    if current and school_year == current:
        candidates.insert(
            0,
            repo_root / "artifacts" / "latest" / "goal_tracking_student_goals" / "raw.xlsx",
        )
    for path in candidates:
        if path.is_file():
            return path.resolve()
    return None


def _goal_progress_worksheet_names(workbook: Any) -> list[str]:
    if GOAL_PROGRESS_SHEET in workbook.sheetnames:
        return [GOAL_PROGRESS_SHEET]
    return list(workbook.sheetnames)


def load_goal_achievement_values_by_student(workbook_path: Path) -> dict[str, set[str]]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        for sheet_name in _goal_progress_worksheet_names(workbook):
            worksheet = workbook[sheet_name]
            header_row: list[str] | None = None
            header_index = -1
            for row_number, row in enumerate(
                worksheet.iter_rows(values_only=True, max_row=10),
                start=1,
            ):
                normalized = [_normalize_header_cell(value) for value in row]
                if "Student ID" in normalized and "Goal Achievement" in normalized:
                    header_row = normalized
                    header_index = row_number
                    break
            if header_row is None:
                continue
            student_index = header_row.index("Student ID")
            achievement_index = header_row.index("Goal Achievement")
            by_student: dict[str, set[str]] = defaultdict(set)
            for row in worksheet.iter_rows(min_row=header_index + 1, values_only=True):
                if not row or not any(row):
                    continue
                if student_index >= len(row) or achievement_index >= len(row):
                    continue
                student_id = clean_value(
                    str(row[student_index]) if row[student_index] is not None else None
                )
                if not student_id:
                    continue
                achievement = clean_value(
                    str(row[achievement_index]) if row[achievement_index] is not None else None
                )
                if achievement:
                    by_student[student_id].add(achievement)
            if by_student:
                return dict(by_student)
        return {}
    finally:
        workbook.close()


def ewgspe_only_student_ids(goal_achievement_by_student: dict[str, set[str]]) -> set[str]:
    excluded: set[str] = set()
    for student_id, values in goal_achievement_by_student.items():
        if not values:
            continue
        if values <= {GOAL_ACHIEVEMENT_EWGSPE}:
            excluded.add(student_id)
    return excluded


def resolve_ewgspe_only_student_ids(
    school_year: str | None,
    *,
    goal_progress_workbook: Path | None = None,
) -> set[str]:
    workbook = goal_progress_workbook or resolve_goal_progress_workbook(school_year)
    if workbook is None:
        return set()
    return ewgspe_only_student_ids(load_goal_achievement_values_by_student(workbook))


def _normalize_school_name(value: str | None) -> str:
    if value is None:
        return ""
    return " ".join(value.lower().replace("&", " and ").split())


def district_key_for_school(school: str | None) -> str | None:
    if not school:
        return None
    if school in SCHOOL_TO_DISTRICT:
        return SCHOOL_TO_DISTRICT[school]
    normalized = _normalize_school_name(school)
    for site, district in SCHOOL_TO_DISTRICT.items():
        if _normalize_school_name(site) == normalized:
            return district
    return None


def is_attendance_only_district(school: str | None) -> bool:
    district = district_key_for_school(school)
    return district in ATTENDANCE_ONLY_DISTRICT_KEYS


def resolve_archived_metrics_workbook(school_year: str) -> Path | None:
    archives_dir = default_archives_dir()
    pull_path = (
        archives_dir
        / school_year
        / "pulls"
        / CISIPHYUS_REPORT_ID
        / "raw.xlsx"
    )
    if pull_path.is_file():
        return pull_path.resolve()
    try:
        src_dir = _default_cisiphyus_root() / "src"
        if str(src_dir) not in sys.path:
            sys.path.insert(0, str(src_dir))
        import config

        programs = config.load_school_year_programs()
        if school_year == config.default_school_year(programs):
            latest_raw = config.latest_raw_path(CISIPHYUS_REPORT_ID)
            if latest_raw.is_file():
                return latest_raw.resolve()
    except (FileNotFoundError, ValueError, OSError):
        pass
    period_dir = archives_dir / school_year / ARCHIVE_PERIOD
    if period_dir.is_dir():
        found = _find_metrics_workbook(period_dir)
        if found is not None:
            return found
    year_dir = archives_dir / school_year
    if year_dir.is_dir():
        for child in sorted(year_dir.iterdir()):
            if child.is_dir():
                found = _find_metrics_workbook(child)
                if found is not None:
                    return found
    return None


def preferred_student_metrics_workbook(
    *,
    local_inputs_dir: Path | None = None,
    school_year: str | None = None,
) -> Path:
    explicit = (os.environ.get("AUDIT_STUDENT_METRICS_WORKBOOK") or "").strip()
    if explicit:
        return Path(explicit).expanduser().resolve()
    base = Path(
        os.environ.get("AUDIT_LOCAL_INPUTS_DIR", "").strip()
        or str(local_inputs_dir or _default_local_inputs_dir())
    ).expanduser().resolve()
    env_name = os.environ.get("AUDIT_STUDENT_METRICS_FILENAME", "").strip()
    if env_name:
        name = env_name
    elif school_year:
        name = student_metrics_workbook_name(school_year)
    else:
        name = default_student_metrics_filename()
    return (base / name).resolve()


def resolve_workbook_for_school_year(
    school_year: str,
    *,
    force_fetch: bool = False,
    fetch_destination: Path | None = None,
) -> Path:
    if not force_fetch:
        archived = resolve_archived_metrics_workbook(school_year)
        if archived is not None:
            print(
                f"[audit] using archived metrics for {school_year}: {archived}",
                file=sys.stderr,
            )
            return archived

    destination = (
        fetch_destination
        or preferred_student_metrics_workbook(school_year=school_year)
    ).resolve()
    fetch_result = fetch_student_metrics_workbook(
        destination=destination,
        force_fetch=force_fetch,
        require_fresh=False,
        school_year=school_year,
    )
    if not fetch_result.succeeded and not fetch_result.destination.exists():
        raise RuntimeError(
            f"no student metrics workbook for {school_year} at {fetch_result.destination}"
        )
    return fetch_result.destination


def workbook_freshness(workbook_path: Path | None) -> dict[str, Any]:
    max_age = float(
        os.environ.get("AUDIT_MAX_AGE_HOURS", "").strip()
        or STUDENT_METRICS_MAX_AGE_HOURS
    )
    freshness: dict[str, Any] = {
        "max_age_hours": max_age,
        "is_stale": True,
        "workbook_exists": False,
    }
    if workbook_path is None or not workbook_path.exists():
        return freshness
    stat = workbook_path.stat()
    age_hours = (datetime.now() - datetime.fromtimestamp(stat.st_mtime)).total_seconds() / 3600
    freshness.update(
        {
            "workbook_exists": True,
            "workbook_age_hours": round(age_hours, 3),
            "workbook_mtime_utc": datetime.fromtimestamp(stat.st_mtime, tz=timezone.utc).isoformat(),
            "is_stale": age_hours > max_age,
        }
    )
    return freshness


def _fetch_reason(*, destination: Path, force_fetch: bool) -> tuple[bool, str]:
    if force_fetch:
        return True, "forced"
    freshness = workbook_freshness(destination if destination.exists() else None)
    if not freshness.get("workbook_exists"):
        return True, "missing"
    if freshness.get("is_stale"):
        return True, "stale"
    return False, "fresh"


def _env_first(*names: str) -> str:
    for name in names:
        value = (os.environ.get(name) or "").strip()
        if value:
            return value
    return ""


def _cisiphyus_cmd(
    cisiphyus_root: Path,
    report_id: str,
    school_year: str | None = None,
) -> list[str]:
    run_py = cisiphyus_root / "run.py"
    cmd: list[str] = [sys.executable, str(run_py), "pull", report_id]
    headed = _flag_true(
        _env_first("AUDIT_CISPHYUS_HEADED", "CISPHYUS_HEADED")
    )
    if headed:
        cmd.append("--headed")
    if school_year:
        cmd.extend(["--school-year", school_year])
    return cmd


def _cisiphyus_raw_path(
    cisiphyus_root: Path,
    report_id: str,
    school_year: str | None = None,
) -> Path:
    if school_year:
        src_dir = cisiphyus_root / "src"
        if str(src_dir) not in sys.path:
            sys.path.insert(0, str(src_dir))
        import config

        programs = config.load_school_year_programs()
        default_sy = config.default_school_year(programs)
        if default_sy and school_year != default_sy:
            return config.archives_raw_path(school_year, report_id)
    return cisiphyus_root / "artifacts" / "latest" / report_id / "raw.xlsx"


def _run_cisiphyus_export(
    *,
    cisiphyus_root: Path,
    report_id: str,
    school_year: str | None = None,
    run: Callable[..., Any],
) -> Path:
    run_py = cisiphyus_root / "run.py"
    if not run_py.is_file():
        raise FileNotFoundError(f"cisiphyus run.py not found at {run_py}")

    completed = run(
        _cisiphyus_cmd(cisiphyus_root, report_id, school_year=school_year),
        cwd=str(cisiphyus_root),
        check=False,
        env=os.environ.copy(),
    )
    if getattr(completed, "returncode", 1) != 0:
        raise RuntimeError(
            f"cisiphyus {report_id} failed with exit code "
            f"{getattr(completed, 'returncode', 'unknown')}"
        )

    raw = _cisiphyus_raw_path(cisiphyus_root, report_id, school_year=school_year)
    if not raw.is_file():
        raise FileNotFoundError(f"expected cisiphyus output missing: {raw}")
    return raw


def fetch_student_metrics_workbook(
    *,
    destination: Path | None = None,
    force_fetch: bool = False,
    require_fresh: bool = False,
    school_year: str | None = None,
    run: Callable[..., Any] = subprocess.run,
) -> FetchResult:
    """Fetch student metrics summary from CISDM when missing, stale, or forced."""
    target = (destination or preferred_student_metrics_workbook()).resolve()
    return fetch_cisiphyus_workbook(
        report_id=CISIPHYUS_REPORT_ID,
        destination=target,
        force_fetch=force_fetch,
        require_fresh=require_fresh,
        school_year=school_year,
        run=run,
    )


def fetch_cisiphyus_workbook(
    *,
    report_id: str,
    destination: Path,
    force_fetch: bool = False,
    require_fresh: bool = False,
    school_year: str | None = None,
    run: Callable[..., Any] = subprocess.run,
) -> FetchResult:
    """Fetch a CISiphyus report export when missing, stale, or forced."""
    target = destination.resolve()
    env_force = _flag_true(os.environ.get("AUDIT_FETCH_FROM_CISDM"))
    should_fetch, reason = _fetch_reason(destination=target, force_fetch=force_fetch or env_force)

    if not should_fetch:
        print(
            f"[audit] workbook fresh at {target} "
            f"(max_age_hours={workbook_freshness(target).get('max_age_hours')}), skipping fetch",
            file=sys.stderr,
        )
        return FetchResult(
            destination=target,
            triggered=False,
            succeeded=True,
            reason="fresh",
        )

    cisiphyus_root = Path(
        _env_first("AUDIT_CISPHYUS_ROOT", "CISIPHYUS_ROOT")
        or str(_default_cisiphyus_root())
    ).expanduser().resolve()

    try:
        raw = _run_cisiphyus_export(
            cisiphyus_root=cisiphyus_root,
            report_id=report_id,
            school_year=school_year,
            run=run,
        )
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(raw, target)

        if require_fresh:
            freshness = workbook_freshness(target)
            if freshness.get("is_stale"):
                age = freshness.get("workbook_age_hours")
                raise RuntimeError(
                    f"stale_workbook_after_refresh: report={report_id} path={target} "
                    f"workbook_age_hours={age} max_age_hours={freshness.get('max_age_hours')}"
                )

        print(
            f"[audit] refreshed {report_id} workbook ({reason}) -> {target}",
            file=sys.stderr,
        )
        return FetchResult(
            destination=target,
            triggered=True,
            succeeded=True,
            reason=reason,
        )
    except (FileNotFoundError, OSError, RuntimeError) as exc:
        if require_fresh and not target.exists():
            raise
        if require_fresh and target.exists():
            freshness = workbook_freshness(target)
            if freshness.get("is_stale"):
                raise RuntimeError(
                    f"stale_workbook_after_failed_refresh: report={report_id} {exc}"
                ) from exc
        print(
            f"[audit] refresh failed ({report_id}, {reason}): {exc}",
            file=sys.stderr,
        )
        if not target.exists():
            raise
        return FetchResult(
            destination=target,
            triggered=True,
            succeeded=False,
            reason=reason,
        )


def fetch_goal_progress_workbook(
    *,
    destination: Path,
    force_fetch: bool = False,
    school_year: str | None = None,
    run: Callable[..., Any] = subprocess.run,
) -> FetchResult:
    """Fetch goal tracking student goals export from CISDM when missing, stale, or forced."""
    return fetch_cisiphyus_workbook(
        report_id=GOAL_TRACKING_REPORT_ID,
        destination=destination,
        force_fetch=force_fetch,
        school_year=school_year,
        run=run,
    )

# --- load workbook ---

HEADER_ROW = 1
DEFAULT_SHEET = "Sheet1"

REQUIRED_HEADERS: tuple[str, ...] = (
    "Organization",
    "School",
    "Student ID",
    "Client ID",
    "Goal",
    "Metric",
    "Baseline",
    "Target",
    "Enrollment Status",
    "School Year",
)

COMPOSITE_KEY_FIELDS = ("Student ID", "School", "School Year", "Goal", "Metric")

GRADING_PERIOD_HEADERS = (
    "1st Grading Period",
    "2nd Grading Period",
    "3rd Grading Period",
    "4th Grading Period",
    "5th Grading Period",
    "6th Grading Period",
)

RecordValue: TypeAlias = str | None | int
RowRecord: TypeAlias = dict[str, RecordValue]


def to_cell_text(value: object) -> str | None:
    if value is None:
        return None
    return str(value)


def row_number(record: RowRecord) -> int:
    value = record.get("row_number")
    if isinstance(value, int):
        return value
    if isinstance(value, str) and value.isdigit():
        return int(value)
    raise KeyError("record missing integer row_number")


def validate_headers(headers: list[str | None], sheet_name: str) -> None:
    present = {h for h in headers if h}
    missing = [h for h in REQUIRED_HEADERS if h not in present]
    if missing:
        raise ValueError(f"Worksheet {sheet_name} is missing required headers: {', '.join(missing)}")


def load_rows(workbook_path: Path, sheet_name: str = DEFAULT_SHEET) -> list[RowRecord]:
    # WHY: data_only matches displayed cell values, not formula strings
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        available = ", ".join(workbook.sheetnames)
        workbook.close()
        raise ValueError(f"Worksheet not found: {sheet_name}. Available: {available}")
    sheet = workbook[sheet_name]
    rows = list(sheet.iter_rows(values_only=True))
    workbook.close()
    if len(rows) < HEADER_ROW:
        raise ValueError(f"Worksheet {sheet_name} has no header row")
    headers = [to_cell_text(v) for v in rows[HEADER_ROW - 1]]
    validate_headers(headers, sheet_name)
    records: list[RowRecord] = []
    for excel_row, row_values in enumerate(rows[HEADER_ROW:], start=HEADER_ROW + 1):
        record: RowRecord = {"row_number": excel_row}
        for index, header in enumerate(headers):
            if not header:
                continue
            cell = row_values[index] if index < len(row_values) else None
            record[header] = to_cell_text(cell)
        records.append(record)
    return records


def load_student_metrics_workbook(
    workbook_path: Path,
    *,
    sheet_name: str = DEFAULT_SHEET,
) -> list[RowRecord]:
    return load_rows(workbook_path, sheet_name=sheet_name)

# --- validity rules ---

Direction = Literal["higher_is_better", "lower_is_better", "none"]

GOAL_DOMAINS: dict[str, str] = {
    "Improve Attendance": "Attendance",
    "Improve School Behavior": "Behavior",
    "Improve Academics": "Academics",
    "Improve Social-Emotional Learning": "SEL",
    "Improve Career Readiness": "Career Readiness",
    "Improve College Readiness": "College Readiness",
    "Improve High Risk Behavior": "High Risk Behavior",
}

ABC_DOMAINS = frozenset({"Attendance", "Behavior", "Academics"})

# ponytail: canonical ABCS goal metric per domain when export has no ABCS flag
PRIMARY_GOAL_METRIC_BY_DOMAIN: dict[str, str] = {
    "Attendance": "Attendance Rate (%)",
    "Behavior": "Suspensions",
    "Academics": "GPA",
}

GOAL_ACHIEVEMENT_EWGSPE = "Exited Within Same Grading Period Enrolled"
GAR_OUTCOME_LABELS = (
    "Goal Met",
    "Goal Not Met, With Progress",
    "Goal Not Met, No Progress",
)
GAR_GOAL_NOT_MET_LABELS = GAR_OUTCOME_LABELS[1:]
GOAL_PROGRESS_SHEET = "CIS_StudentProgress_Detail"
GOAL_PROGRESS_HEADER_ROW = 3
GOAL_PROGRESS_FILENAME_TEMPLATE = "{school_year}_GoalAchievement.xlsx"

# ponytail: attendance-only ABCS scope for these district keys (Jon Irons, SY25-26)
ATTENDANCE_ONLY_DISTRICT_KEYS = frozenset({"Pocono", "SE Delco"})

SCHOOL_TO_DISTRICT: dict[str, str] = {
    "Octorara Junior and Senior High School": "Octorara",
    "Freemansburg Elementary School": "Bethlehem",
    "10th & Penn": "Reading",
    "Stony Creek Elementary": "Antietam",
    "Tyson-Schoener Elementary": "Reading",
    "Broughal Middle School": "Bethlehem",
    "Fountain Hill Elementary School": "Bethlehem",
    "Mosser Elementary School": "Allentown",
    "Sharon Hill School": "SE Delco",
    "Delcroft School": "SE Delco",
    "Southwest Middle School": "Reading",
    "Darby Township School": "SE Delco",
    "Wilson Area Intermediate School": "Wilson Area",
    "Northwestern Lehigh High School": "Northwestern Lehigh",
    "Fleetwood Area Middle School": "Fleetwood",
    "Southern Middle School": "Reading",
    "Clear Run Elementary Center": "Pocono",
    "Pocono Mountain West Junior High School": "Pocono",
    "Harrison Morton Middle School": "Allentown",
    "Octorara Intermediate School": "Octorara",
    "Harris School": "SE Delco",
    "Northeast Middle School": "Bethlehem",
    "Fleetwood High School": "Fleetwood",
    "Trexler Middle School": "Allentown",
    "Eyer Middle School": "East Penn",
    "Clear Run Intermediate School": "Pocono",
    "South Mountain Middle School": "Allentown",
    "Lincoln Leadership Academy": "LLA",
    "Whitehall-Coplay Middle School": "Whitehall",
    "Academy Park High School": "SE Delco",
    "Avon Grove Intermediate School": "Avon Grove",
    "Pocono Mountain West High School": "Pocono",
    "Avon Grove High School": "Avon Grove",
    "Freedom High School": "Bethlehem",
    "Easton Area Middle School": "Easton",
    "Lehigh Career & Technical Institute": "LCTI",
    "Liberty High School": "Bethlehem",
    "Easton Area High School": "Easton",
    "Emmaus High School": "East Penn",
    "William Allen High School": "Allentown",
    "Reading High School": "Reading",
}

STUDENT_COMPLETENESS_ISSUE_CODES = frozenset(
    {
        "abc_baselines_incomplete",
        "abc_target_missing",
        "attendance_baseline_missing",
        "attendance_target_missing",
    }
)

METRIC_DOMAIN_OVERRIDES: dict[str, frozenset[str]] = {
    "Attendance Rate (%)": frozenset({"Attendance"}),
    "Attendance Rate (days absent)": frozenset({"Attendance"}),
    "Tardies": frozenset({"Attendance", "Behavior"}),
    "Suspensions": frozenset({"Behavior"}),
    "Disciplinary Referrals": frozenset({"Behavior"}),
    "Conduct": frozenset({"Behavior"}),
    "Other Behavior Incidents": frozenset({"Behavior"}),
    "Credits Needed/Credit Completion": frozenset({"Academics"}),
    "Reading Level": frozenset({"Academics"}),
    "GPA": frozenset({"Academics"}),
    "Overall SEAD assessment score": frozenset({"SEL"}),
    "Self Control (SEAD)": frozenset({"SEL"}),
    "Social Awareness (SEAD)": frozenset({"SEL"}),
    "Self-Perception (SEAD)": frozenset({"SEL"}),
    "Social Support (SEAD)": frozenset({"SEL"}),
    "Global Engagement (SEAD)": frozenset({"SEL"}),
    "Social Engagement": frozenset({"SEL"}),
    "Behavior Engagement": frozenset({"SEL"}),
    "Emotional Engagement": frozenset({"SEL"}),
    "Global Engagement (Engagement Survey)": frozenset({"SEL"}),
    "Cognitive Engagement": frozenset({"SEL"}),
    "Other (SEL)": frozenset({"SEL"}),
    "Other High Risk Behavior": frozenset({"High Risk Behavior"}),
    "Drank Alcohol (# of times past 30 days)": frozenset({"High Risk Behavior"}),
    "Violent/bullied someone (# of times past 30 days)": frozenset({"High Risk Behavior"}),
    "CTE Completion": frozenset({"Academics"}),
    "Complete career assessment": frozenset({"Career Readiness"}),
    "Accept a position of employment": frozenset({"Career Readiness"}),
    "Interview with one or more potential employers": frozenset({"Career Readiness"}),
    "Create a resume": frozenset({"Career Readiness"}),
    "Other Career Readiness": frozenset({"Career Readiness"}),
    "Apply to one or more colleges/universities": frozenset({"College Readiness"}),
    "Accepted to one or more colleges/universities": frozenset({"College Readiness"}),
    "Other College Readiness": frozenset({"College Readiness"}),
}

# WHY: Improve Attendance goals may allow one unit of slack (e.g. one extra tardy)
ATTENDANCE_GOAL_PERMISSIVE_METRICS = frozenset(
    {
        "Attendance Rate (%)",
        "Tardies",
    }
)
ATTENDANCE_GOAL_PERMISSIVE_SLACK = 1.0


def attendance_goal_permissive_slack(metric: str | None, goal: str | None) -> float:
    if goal_domain(goal) != "Attendance":
        return 0.0
    metric_name = clean_value(metric)
    if metric_name in ATTENDANCE_GOAL_PERMISSIVE_METRICS:
        return ATTENDANCE_GOAL_PERMISSIVE_SLACK
    return 0.0


GRADE_SCALE: frozenset[str] = frozenset(
    {"F", "D-", "D", "D+", "C-", "C", "C+", "B-", "B", "B+", "A-", "A", "A+"}
)

GRADE_ORDER: tuple[str, ...] = (
    "F",
    "D-",
    "D",
    "D+",
    "C-",
    "C",
    "C+",
    "B-",
    "B",
    "B+",
    "A-",
    "A",
    "A+",
)

GRADE_RANK: dict[str, int] = {grade: index for index, grade in enumerate(GRADE_ORDER)}


READING_LEVEL_TERMS: frozenset[str] = frozenset(
    {
        "below",
        "below benchmark",
        "well below benchmark",
        "well-below benchmark",
        "at benchmark",
        "at or above benchmark",
        "above benchmark",
        "at grade level",
        "below grade level",
        "expected reading level",
        "expected grade reading level",
        "at or above grade level",
    }
)

READING_LEVEL_NORMALIZATION: dict[str, str] = {
    "well-below benchmark": "well below benchmark",
    "below reading level": "below grade level",
    "expected grade reading level": "expected reading level",
    "proficient": "at benchmark",
    "at risk": "below benchmark",
}

ENGAGEMENT_TERMS: frozenset[str] = frozenset(
    {
        "lower engagement",
        "low engagement",
        "moderate engagement",
        "higher engagement",
        "high engagement",
    }
)

BOOLEAN_TERMS: frozenset[str] = frozenset({"yes", "no"})

ANNOTATION_EXEMPT_SCALES: frozenset[str] = frozenset(
    {"reading_level", "engagement_text", "numeric_or_engagement", "boolean"}
)

PAREN_ANNOTATION_RE = re.compile(r"\([^)]+\)")


def clean_value(value: str | None) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def parse_number(value: str | None) -> float | None:
    text = clean_value(value)
    if text is None:
        return None
    cleaned = text.replace(",", "").replace("%", "")
    try:
        return float(cleaned)
    except ValueError:
        return None


def normalize_text(value: str | None) -> str | None:
    text = clean_value(value)
    if text is None:
        return None
    return text.lower()


def normalize_grade(value: str | None) -> str | None:
    text = clean_value(value)
    if text is None:
        return None
    return text.split("(")[0].strip().upper()


def normalize_reading_level(value: str | None) -> str | None:
    text = normalize_text(value)
    if text is None:
        return None
    canonical = " ".join(text.replace("-", " ").split())
    if canonical.startswith("level "):
        suffix = canonical.removeprefix("level ").strip()
        if parse_number(suffix) is not None:
            return suffix
    if canonical.endswith(" letters"):
        prefix = canonical.removesuffix(" letters").strip()
        if parse_number(prefix) is not None:
            return prefix
    return READING_LEVEL_NORMALIZATION.get(canonical, canonical)


def goal_domain(goal: str | None) -> str | None:
    return GOAL_DOMAINS.get(clean_value(goal) or "")


def metric_domains(metric: str | None) -> frozenset[str]:
    metric_name = clean_value(metric)
    if metric_name is None:
        return frozenset()
    if metric_name in METRIC_DOMAIN_OVERRIDES:
        return METRIC_DOMAIN_OVERRIDES[metric_name]
    lowered = metric_name.lower()
    if "attendance rate" in lowered:
        return frozenset({"Attendance"})
    if "suspension" in lowered or "referral" in lowered or "incident" in lowered or "conduct" in lowered:
        return frozenset({"Behavior"})
    if "grade" in lowered or "gpa" in lowered or "reading" in lowered or "credit" in lowered or "score" in lowered:
        return frozenset({"Academics"})
    if "sead" in lowered or "engagement" in lowered or "(sel)" in lowered:
        return frozenset({"SEL"})
    if "career" in lowered or "resume" in lowered or "employment" in lowered:
        return frozenset({"Career Readiness"})
    if "college" in lowered or "universit" in lowered:
        return frozenset({"College Readiness"})
    if (
        "tobacco" in lowered
        or "alcohol" in lowered
        or "high risk" in lowered
        or "bullied" in lowered
        or "violent" in lowered
    ):
        return frozenset({"High Risk Behavior"})
    return frozenset()


def metric_scale(metric: str | None) -> str | None:
    metric_name = clean_value(metric)
    if metric_name is None:
        return None
    lowered = metric_name.lower()
    if metric_name == "Attendance Rate (%)":
        return "percent"
    if (
        "days absent" in lowered
        or "suspension" in lowered
        or "referral" in lowered
        or "incident" in lowered
        or "tard" in lowered
        or "conduct" in lowered
    ):
        return "count"
    if "core course grades" in lowered:
        return "grade_or_numeric"
    if metric_name == "GPA" or "score" in lowered:
        return "numeric"
    if metric_name == "Reading Level":
        return "reading_level"
    if "credits needed" in lowered or "credit completion" in lowered:
        return "numeric"
    if metric_name == "Other (SEL)":
        return "numeric_or_engagement"
    if metric_name in {
        "Complete career assessment",
        "Accept a position of employment",
        "Interview with one or more potential employers",
        "Create a resume",
        "Apply to one or more colleges/universities",
        "Accepted to one or more colleges/universities",
        "CTE Completion",
    }:
        return "boolean"
    if metric_name in {"Other Career Readiness", "Other College Readiness"}:
        return "numeric"
    if "emotional engagement" in lowered:
        return "engagement_text"
    if "engagement" in lowered or "sead" in lowered:
        return "numeric_or_engagement"
    if "alcohol" in lowered or "high risk" in lowered or "bullied" in lowered or "violent" in lowered:
        return "count"
    return None


def metric_direction(metric: str | None) -> Direction:
    metric_name = clean_value(metric)
    # WHY: CISDM mixes rate and days-absent semantics; direction check is unreliable here
    if metric_name == "Attendance Rate (days absent)":
        return "none"
    scale = metric_scale(metric)
    if scale in {"boolean", "engagement_text", "reading_level"}:
        return "none"
    if scale == "numeric_or_engagement":
        return "higher_is_better"
    if scale == "count":
        return "lower_is_better"
    if scale == "percent":
        return "higher_is_better"
    if scale == "numeric":
        return "higher_is_better"
    if scale == "grade_or_numeric":
        return "higher_is_better"
    return "none"


def value_matches_scale(metric: str | None, value: str | None) -> bool:
    text = clean_value(value)
    if text is None:
        return True
    scale = metric_scale(metric)
    if scale is None:
        return False
    lowered = normalize_text(text) or ""
    if scale == "percent":
        return parse_number(text) is not None and "day" not in lowered and normalize_grade(text) not in GRADE_SCALE
    if scale == "count":
        return parse_number(text) is not None
    if scale == "numeric":
        return parse_number(text) is not None
    if scale == "grade_or_numeric":
        return parse_number(text) is not None or normalize_grade(text) in GRADE_SCALE
    if scale == "reading_level":
        normalized = normalize_reading_level(text)
        return normalized in READING_LEVEL_TERMS or parse_number(normalized) is not None
    if scale == "engagement_text":
        return lowered in ENGAGEMENT_TERMS
    if scale == "numeric_or_engagement":
        return parse_number(text) is not None or lowered in ENGAGEMENT_TERMS
    if scale == "boolean":
        return lowered in BOOLEAN_TERMS
    return False


def has_undocumented_annotation(metric: str | None, value: str | None) -> bool:
    text = clean_value(value)
    if text is None:
        return False
    scale = metric_scale(metric)
    if scale is None or scale in ANNOTATION_EXEMPT_SCALES:
        return False
    return bool(PAREN_ANNOTATION_RE.search(text))


def _numeric_direction_ok(
    baseline: float,
    target: float,
    direction: Direction,
    *,
    permissive_slack: float = 0.0,
) -> bool:
    if direction == "higher_is_better":
        return target >= baseline - permissive_slack
    if direction == "lower_is_better":
        return target <= baseline + permissive_slack
    return True


def _grade_direction_ok(
    baseline_grade: str,
    target_grade: str,
    direction: Direction,
) -> bool:
    baseline_rank = GRADE_RANK.get(baseline_grade)
    target_rank = GRADE_RANK.get(target_grade)
    if baseline_rank is None or target_rank is None:
        return True
    if direction == "higher_is_better":
        return target_rank >= baseline_rank
    if direction == "lower_is_better":
        return target_rank <= baseline_rank
    return True


def baseline_target_direction_mismatch(
    metric: str | None,
    baseline: str | None,
    target: str | None,
    goal: str | None = None,
) -> bool:
    direction = metric_direction(metric)
    if direction == "none":
        return False
    baseline_text = clean_value(baseline)
    target_text = clean_value(target)
    if baseline_text is None or target_text is None:
        return False

    baseline_num = parse_number(baseline_text)
    target_num = parse_number(target_text)
    if baseline_num is not None and target_num is not None:
        slack = attendance_goal_permissive_slack(metric, goal)
        return not _numeric_direction_ok(
            baseline_num,
            target_num,
            direction,
            permissive_slack=slack,
        )

    baseline_grade = normalize_grade(baseline_text)
    target_grade = normalize_grade(target_text)
    if (
        baseline_grade in GRADE_SCALE
        and target_grade in GRADE_SCALE
        and baseline_grade is not None
        and target_grade is not None
    ):
        return not _grade_direction_ok(baseline_grade, target_grade, direction)
    return False


def analyze_record(record: RowRecord) -> dict[str, object]:
    goal = record.get("Goal")
    metric = record.get("Metric")
    baseline = record.get("Baseline")
    target = record.get("Target")
    issue_codes: list[str] = []

    expected_domain = goal_domain(str(goal) if goal is not None else None)
    allowed_domains = metric_domains(str(metric) if metric is not None else None)
    if expected_domain and allowed_domains and expected_domain not in allowed_domains:
        issue_codes.append("goal_metric_domain_mismatch")
    elif expected_domain is None or not allowed_domains:
        issue_codes.append("manual_review_domain")

    metric_str = str(metric) if metric is not None else None
    baseline_str = str(baseline) if baseline is not None else None
    target_str = str(target) if target is not None else None

    if not value_matches_scale(metric_str, baseline_str):
        issue_codes.append("baseline_scale_mismatch")
    if not value_matches_scale(metric_str, target_str):
        issue_codes.append("target_scale_mismatch")

    if has_undocumented_annotation(metric_str, baseline_str):
        issue_codes.append("undocumented_value_annotation")
    if has_undocumented_annotation(metric_str, target_str):
        issue_codes.append("undocumented_value_annotation")

    goal_str = str(goal) if goal is not None else None
    if baseline_target_direction_mismatch(metric_str, baseline_str, target_str, goal=goal_str):
        issue_codes.append("baseline_target_direction_mismatch")

    return {
        "school": clean_value(str(record.get("School")) if record.get("School") is not None else None),
        "goal": clean_value(str(goal) if goal is not None else None),
        "metric": clean_value(metric_str),
        "expected_domain": expected_domain,
        "allowed_domains": sorted(allowed_domains),
        "issue_codes": issue_codes,
    }

# --- run checks and write outputs ---

FLAG_DESCRIPTIONS: dict[str, str] = {
    "goal_metric_domain_mismatch": "Goal domain does not match metric",
    "manual_review_domain": "Goal or metric domain needs manual review",
    "baseline_scale_mismatch": "Baseline value does not match metric scale",
    "target_scale_mismatch": "Target value does not match metric scale",
    "baseline_target_direction_mismatch": "Target is not better than baseline for metric family",
    "undocumented_value_annotation": "Value contains undocumented parenthetical annotation",
    "both_baseline_and_target_blank": "Both baseline and target are blank",
    "both_blank_with_two_progress_reports": "Both blank but two or more grading periods filled",
    "target_without_baseline": "Target set without baseline",
    "baseline_without_target": "Baseline set without target",
    "baseline_without_target_supplemental_metric": (
        "Baseline without target (non-goal metric — student has another complete goal)"
    ),
    "target_without_baseline_supplemental_metric": (
        "Target without baseline (non-goal metric — student has another complete goal)"
    ),
    "baseline_without_target_domain_tracking": (
        "Baseline without target (supplemental metric — domain goal already set)"
    ),
    "baseline_without_target_no_goal_context": (
        "Baseline without target (no ABC goals set yet — accepted)"
    ),
    "target_without_baseline_domain_tracking": (
        "Target without baseline (supplemental metric — domain goal already set)"
    ),
    "target_without_baseline_no_goal_context": (
        "Target without baseline (no ABC goals set yet — accepted)"
    ),
    "duplicate_composite_key": "Duplicate Student ID+School+Goal+Metric row",
    "student_client_id_mismatch": "Student ID maps to multiple Client IDs",
    "case_manager_blank": "Case Manager is blank",
    "abc_baselines_incomplete": "Student missing one or more ABC goal baselines",
    "abc_target_missing": "Student missing target on all ABC goals",
    "attendance_baseline_missing": "Attendance-only district: attendance baseline missing",
    "attendance_target_missing": "Attendance-only district: attendance target missing",
}

ACCEPTED_ISSUE_CODES = frozenset(
    {
        "baseline_without_target_no_goal_context",
        "baseline_without_target_domain_tracking",
        "baseline_without_target_supplemental_metric",
        "target_without_baseline_no_goal_context",
        "target_without_baseline_domain_tracking",
        "target_without_baseline_supplemental_metric",
    }
)

PARTIAL_METRIC_ISSUE_CODES = frozenset(
    {
        "baseline_without_target",
        "target_without_baseline",
    }
)

BASELINE_SUPPLEMENTAL_ACCEPTED_CODES = frozenset(
    {
        "baseline_without_target_domain_tracking",
        "baseline_without_target_supplemental_metric",
    }
)

TARGET_SUPPLEMENTAL_ACCEPTED_CODES = frozenset(
    {
        "target_without_baseline_domain_tracking",
        "target_without_baseline_supplemental_metric",
    }
)

NO_GOAL_CONTEXT_ACCEPTED_CODES = frozenset(
    {
        "baseline_without_target_no_goal_context",
        "target_without_baseline_no_goal_context",
    }
)

DOMAIN_TRACKING_ACCEPTED_CODES = frozenset(
    {
        "baseline_without_target_domain_tracking",
        "target_without_baseline_domain_tracking",
    }
)

SUPPLEMENTAL_ACCEPTED_CODES = frozenset(
    {
        "baseline_without_target_supplemental_metric",
        "target_without_baseline_supplemental_metric",
    }
)

ACCEPTED_EXPORT_SPECS: tuple[tuple[str, frozenset[str], str], ...] = (
    ("accepted_no_goal_context", NO_GOAL_CONTEXT_ACCEPTED_CODES, "accepted_no_goal_context.csv"),
    ("accepted_domain_tracking", DOMAIN_TRACKING_ACCEPTED_CODES, "accepted_domain_tracking.csv"),
    ("accepted_supplemental", SUPPLEMENTAL_ACCEPTED_CODES, "accepted_supplemental.csv"),
)


@dataclass
class AuditConfig:
    sheet_name: str = DEFAULT_SHEET
    include_ok_rows: bool = False
    school_year: str | None = None
    goal_progress_workbook: Path | None = None
    excluded_student_ids: set[str] | None = None


def _text(record: RowRecord, field: str) -> str | None:
    value = record.get(field)
    return value if isinstance(value, str) else None


def is_blank(value: str | None) -> bool:
    return clean_value(value) is None


def baseline_target_status(record: RowRecord) -> str:
    baseline_blank = is_blank(_text(record, "Baseline"))
    target_blank = is_blank(_text(record, "Target"))
    if baseline_blank and target_blank:
        return "both_blank"
    if baseline_blank and not target_blank:
        return "target_only"
    if not baseline_blank and target_blank:
        return "baseline_only"
    return "both_present"


ProgressStatus = Literal["on_track", "off_track", "no_progress_data", "indeterminate"]
PROGRESS_STATUS_ORDER = ("on_track", "off_track", "no_progress_data", "indeterminate")


def latest_metric_value(record: RowRecord) -> str | None:
    latest = clean_value(_text(record, "Latest Progress"))
    if latest is not None:
        return latest
    return latest_grading_period_value(record)


def latest_grading_period_value(record: RowRecord) -> str | None:
    for header in reversed(GRADING_PERIOD_HEADERS):
        text = _text(record, header)
        if not is_blank(text):
            return clean_value(text)
    return None


def progress_meets_target(
    metric: str | None,
    progress: str | None,
    target: str | None,
) -> bool | None:
    direction = metric_direction(metric)
    if direction == "none":
        return None
    progress_text = clean_value(progress)
    target_text = clean_value(target)
    if progress_text is None or target_text is None:
        return None

    progress_num = parse_number(progress_text)
    target_num = parse_number(target_text)
    if progress_num is not None and target_num is not None:
        if direction == "higher_is_better":
            return progress_num >= target_num
        if direction == "lower_is_better":
            return progress_num <= target_num
        return None

    progress_grade = normalize_grade(progress_text)
    target_grade = normalize_grade(target_text)
    if (
        progress_grade in GRADE_SCALE
        and target_grade in GRADE_SCALE
        and progress_grade is not None
        and target_grade is not None
    ):
        progress_rank = GRADE_RANK.get(progress_grade)
        target_rank = GRADE_RANK.get(target_grade)
        if progress_rank is None or target_rank is None:
            return None
        if direction == "higher_is_better":
            return progress_rank >= target_rank
        if direction == "lower_is_better":
            return progress_rank <= target_rank
    return None


def progress_status_for_record(record: RowRecord) -> ProgressStatus:
    if baseline_target_status(record) != "both_present":
        return "no_progress_data"
    progress = latest_grading_period_value(record)
    if progress is None:
        return "no_progress_data"
    metric = _text(record, "Metric")
    if not value_matches_scale(metric, progress):
        return "indeterminate"
    meets = progress_meets_target(metric, progress, _text(record, "Target"))
    if meets is None:
        return "indeterminate"
    return "on_track" if meets else "off_track"


def progress_row_key(record: RowRecord) -> str:
    return "|".join(
        clean_value(_text(record, field)) or ""
        for field in ("Student ID", "School", "Goal", "Metric")
    )


def _empty_progress_counts() -> dict[str, int]:
    return {status: 0 for status in PROGRESS_STATUS_ORDER}


def progress_rollup_from_records(records: list[RowRecord]) -> dict[str, Any]:
    global_counts: Counter[str] = Counter()
    schools: dict[str, Counter[str]] = defaultdict(Counter)
    progress_index: dict[str, str] = {}
    eligible_rows = 0

    for record in records:
        if baseline_target_status(record) != "both_present":
            continue
        eligible_rows += 1
        status = progress_status_for_record(record)
        global_counts[status] += 1
        school = clean_value(_text(record, "School")) or "(unknown)"
        schools[school][status] += 1
        progress_index[progress_row_key(record)] = status

    return {
        "global": {status: int(global_counts.get(status, 0)) for status in PROGRESS_STATUS_ORDER},
        "schools": {
            school: {status: int(counts.get(status, 0)) for status in PROGRESS_STATUS_ORDER}
            for school, counts in sorted(schools.items())
        },
        "eligible_rows": eligible_rows,
        "progress_index": progress_index,
    }


def progress_rollup(
    workbook_path: Path,
    sheet_name: str = DEFAULT_SHEET,
) -> dict[str, Any]:
    return progress_rollup_from_records(load_rows(workbook_path, sheet_name))


def grading_period_count(record: RowRecord) -> int:
    return sum(1 for header in GRADING_PERIOD_HEADERS if not is_blank(_text(record, header)))


def grading_period_fill_stats(
    workbook_path: Path,
    sheet_name: str = DEFAULT_SHEET,
) -> dict[str, int]:
    rows = load_rows(workbook_path, sheet_name)
    stats = {header: 0 for header in GRADING_PERIOD_HEADERS}
    for record in rows:
        for header in GRADING_PERIOD_HEADERS:
            if not is_blank(_text(record, header)):
                stats[header] += 1
    return stats


def student_metric_context(record: RowRecord) -> tuple[str, str]:
    return (
        clean_value(_text(record, "Student ID")) or "",
        clean_value(_text(record, "School Year")) or "",
    )


EXCLUDED_ENROLLMENT_STATUSES = frozenset({"exited"})


def filter_records_excluding_students(
    records: list[RowRecord],
    excluded_student_ids: set[str],
) -> tuple[list[RowRecord], int]:
    if not excluded_student_ids:
        return records, 0
    kept: list[RowRecord] = []
    excluded_rows = 0
    for record in records:
        student_id = clean_value(_text(record, "Student ID")) or ""
        if student_id in excluded_student_ids:
            excluded_rows += 1
        else:
            kept.append(record)
    return kept, excluded_rows


def student_abc_domains_with_baseline(
    records: list[RowRecord],
    *,
    student_id: str,
    school_year: str,
) -> set[str]:
    domains: set[str] = set()
    for record in records:
        sid, year = student_metric_context(record)
        if sid != student_id or year != school_year:
            continue
        domain = record_abc_domain(record)
        if domain is None or is_blank(_text(record, "Baseline")):
            continue
        domains.add(domain)
    return domains


def student_abc_domains_with_target(
    records: list[RowRecord],
    *,
    student_id: str,
    school_year: str,
    domains: frozenset[str],
) -> set[str]:
    with_target: set[str] = set()
    for record in records:
        sid, year = student_metric_context(record)
        if sid != student_id or year != school_year:
            continue
        domain = record_abc_domain(record)
        if domain is None or domain not in domains or is_blank(_text(record, "Target")):
            continue
        with_target.add(domain)
    return with_target


def _group_records_by_student(records: list[RowRecord]) -> dict[tuple[str, str], list[RowRecord]]:
    grouped: dict[tuple[str, str], list[RowRecord]] = defaultdict(list)
    for record in records:
        student_id, school_year = student_metric_context(record)
        if student_id:
            grouped[(student_id, school_year)].append(record)
    return grouped


def _audit_student_abc_completeness(
    records: list[RowRecord],
    *,
    issue_codes: Counter[str],
    detail_map: dict[int, list[str]],
) -> int:
    flagged_students = 0
    for (student_id, school_year), student_records in _group_records_by_student(records).items():
        del student_id, school_year
        anchor = min(student_records, key=row_number)
        school = clean_value(_text(anchor, "School"))
        attendance_only = is_attendance_only_district(school)
        required_domains = (
            frozenset({"Attendance"}) if attendance_only else ABC_DOMAINS
        )
        sid, year = student_metric_context(anchor)
        baselines = student_abc_domains_with_baseline(
            records,
            student_id=sid,
            school_year=year,
        )
        targets = student_abc_domains_with_target(
            records,
            student_id=sid,
            school_year=year,
            domains=required_domains,
        )
        missing_baselines = required_domains - baselines
        codes: list[str] = []
        if missing_baselines:
            if attendance_only:
                codes.append("attendance_baseline_missing")
            else:
                codes.append("abc_baselines_incomplete")
        if not targets:
            if attendance_only:
                codes.append("attendance_target_missing")
            else:
                codes.append("abc_target_missing")
        if codes:
            for code in codes:
                issue_codes[code] += 1
                detail_map[row_number(anchor)].append(code)
            flagged_students += 1
    return flagged_students


def is_audit_eligible_record(
    record: RowRecord,
    *,
    final_goal_achievement_students: set[str] | None = None,
) -> bool:
    status = clean_value(_text(record, "Enrollment Status"))
    if not status or status.lower() not in EXCLUDED_ENROLLMENT_STATUSES:
        return True
    if final_goal_achievement_students is None:
        return False
    student_id = clean_value(_text(record, "Student ID")) or ""
    return student_id in final_goal_achievement_students


def filter_audit_records(
    records: list[RowRecord],
    *,
    final_goal_achievement_students: set[str] | None = None,
    apply_exited_filter: bool = True,
) -> tuple[list[RowRecord], int, int]:
    if not apply_exited_filter:
        return records, 0, 0
    kept: list[RowRecord] = []
    excluded = 0
    included_exited = 0
    for record in records:
        if is_audit_eligible_record(
            record,
            final_goal_achievement_students=final_goal_achievement_students,
        ):
            kept.append(record)
            status = clean_value(_text(record, "Enrollment Status"))
            if status and status.lower() in EXCLUDED_ENROLLMENT_STATUSES:
                included_exited += 1
        else:
            excluded += 1
    return kept, excluded, included_exited


def is_closed_school_year(school_year: str | None) -> bool:
    if not school_year:
        return False
    current = current_school_year()
    return current is not None and school_year != current


def enrollment_filter_plan(
    school_year: str | None,
    *,
    final_goal_achievement_students: set[str] | None,
) -> tuple[bool, set[str] | None]:
    # WHY: closed-year re-pulls mark nearly everyone Exited; year export is the scope
    if is_closed_school_year(school_year):
        return False, None
    if final_goal_achievement_students is not None:
        return True, final_goal_achievement_students
    if should_exclude_exited_students(school_year):
        return True, None
    return False, None


def should_exclude_exited_students(school_year: str | None) -> bool:
    if not school_year:
        return False
    current = current_school_year()
    return current is not None and school_year == current


def record_abc_domain(record: RowRecord) -> str | None:
    goal = clean_value(_text(record, "Goal"))
    if goal is None:
        return None
    domain = goal_domain(goal)
    return domain if domain in ABC_DOMAINS else None


def student_domains_with_complete_metric(records: list[RowRecord]) -> set[tuple[str, str, str]]:
    complete: set[tuple[str, str, str]] = set()
    for record in records:
        domain = record_abc_domain(record)
        if domain is None:
            continue
        student_id, school_year = student_metric_context(record)
        if (
            student_id
            and not is_blank(_text(record, "Baseline"))
            and not is_blank(_text(record, "Target"))
        ):
            complete.add((student_id, school_year, domain))
    return complete


def students_with_any_abc_complete(records: list[RowRecord]) -> set[tuple[str, str]]:
    students: set[tuple[str, str]] = set()
    for student_id, school_year, _domain in student_domains_with_complete_metric(records):
        students.add((student_id, school_year))
    return students


def is_primary_goal_metric(record: RowRecord) -> bool:
    domain = record_abc_domain(record)
    if domain is None:
        return False
    primary = PRIMARY_GOAL_METRIC_BY_DOMAIN.get(domain)
    metric = clean_value(_text(record, "Metric"))
    return bool(primary and metric == primary)


def _partial_metric_disposition(
    record: RowRecord,
    *,
    status: str,
    complete_student_contexts: set[tuple[str, str]],
) -> tuple[str | None, str | None]:
    """Return (issue_code, accepted_code) for baseline_only / target_only rows."""
    student_id, school_year = student_metric_context(record)
    if not student_id:
        return None, None

    issue_code = "baseline_without_target" if status == "baseline_only" else "target_without_baseline"
    if (student_id, school_year) in complete_student_contexts:
        return None, f"{issue_code}_supplemental_metric"
    return issue_code, None


def student_contexts_with_complete_metric(records: list[RowRecord]) -> set[tuple[str, str]]:
    complete: set[tuple[str, str]] = set()
    for record in records:
        context = student_metric_context(record)
        if context[0] and not is_blank(_text(record, "Baseline")) and not is_blank(_text(record, "Target")):
            complete.add(context)
    return complete


class IssueSink:
    def __init__(
        self,
        *,
        issue_codes: Counter[str],
        detail_map: dict[int, list[str]],
    ) -> None:
        self._issue_codes = issue_codes
        self._detail_map = detail_map

    def add(self, row_num: int, code: str) -> None:
        self._issue_codes[code] += 1
        self._detail_map[row_num].append(code)

    def add_many(self, row_num: int, codes: list[str]) -> None:
        for code in codes:
            self.add(row_num, code)


class AcceptedSink:
    def __init__(
        self,
        *,
        accepted_codes: Counter[str],
        detail_map: dict[int, list[str]],
    ) -> None:
        self._accepted_codes = accepted_codes
        self._detail_map = detail_map

    def add(self, row_num: int, code: str) -> None:
        self._accepted_codes[code] += 1
        self._detail_map[row_num].append(code)

    def add_many(self, row_num: int, codes: list[str]) -> None:
        for code in codes:
            self.add(row_num, code)


def audit_composite_keys(records: list[RowRecord]) -> dict[str, Any]:
    seen: dict[tuple[str, str, str, str, str], int] = {}
    duplicates: list[int] = []
    for record in records:
        key = tuple(clean_value(_text(record, f)) or "" for f in COMPOSITE_KEY_FIELDS)
        if key in seen:
            duplicates.append(row_number(record))
        else:
            seen[key] = row_number(record)
    return {
        "unique_keys": len(seen),
        "duplicate_row_numbers": duplicates,
        "duplicate_count": len(duplicates),
    }


def audit_student_client_ids(records: list[RowRecord]) -> dict[str, Any]:
    sid_to_cids: dict[str, set[str]] = defaultdict(set)
    for record in records:
        sid = clean_value(_text(record, "Student ID"))
        cid = clean_value(_text(record, "Client ID"))
        if sid and cid:
            sid_to_cids[sid].add(cid)
    conflicting = {sid for sid, cids in sid_to_cids.items() if len(cids) > 1}
    return {
        "student_ids_with_multiple_client_ids": len(conflicting),
        "conflicting_student_ids": sorted(conflicting),
    }


def audit_baseline_target(records: list[RowRecord]) -> dict[str, int]:
    counts: Counter[str] = Counter()
    for record in records:
        status = baseline_target_status(record)
        if status == "both_blank":
            counts["both_baseline_and_target_blank"] += 1
        elif status == "target_only":
            counts["target_without_baseline"] += 1
        elif status == "baseline_only":
            counts["baseline_without_target"] += 1
        else:
            counts["baseline_and_target_present"] += 1
    return dict(counts)


def _audit_case_manager_blanks(records: list[RowRecord], sink: IssueSink) -> None:
    if not any("Case Manager" in record for record in records):
        return
    for record in records:
        if is_blank(_text(record, "Case Manager")):
            sink.add(row_number(record), "case_manager_blank")


def _mark_duplicate_composite_keys(key_info: dict[str, Any], sink: IssueSink) -> None:
    for row_num in key_info.get("duplicate_row_numbers") or []:
        sink.add(row_num, "duplicate_composite_key")


def _mark_student_client_id_mismatches(
    records: list[RowRecord],
    *,
    conflict_sid_set: set[str],
    sink: IssueSink,
) -> None:
    if not conflict_sid_set:
        return
    for record in records:
        sid = clean_value(_text(record, "Student ID"))
        if sid and sid in conflict_sid_set:
            sink.add(row_number(record), "student_client_id_mismatch")


def _audit_baseline_target_flags(
    records: list[RowRecord],
    *,
    complete_student_contexts: set[tuple[str, str]],
    sink: IssueSink,
    accepted_sink: AcceptedSink,
) -> None:
    for record in records:
        status = baseline_target_status(record)
        flags: list[str] = []
        accepted_flags: list[str] = []
        if status == "both_blank":
            flags.append("both_baseline_and_target_blank")
            if grading_period_count(record) >= 2:
                flags.append("both_blank_with_two_progress_reports")
        elif status in ("target_only", "baseline_only"):
            issue_code, accepted_code = _partial_metric_disposition(
                record,
                status=status,
                complete_student_contexts=complete_student_contexts,
            )
            if issue_code:
                flags.append(issue_code)
            if accepted_code:
                accepted_flags.append(accepted_code)
        if flags:
            sink.add_many(row_number(record), flags)
        if accepted_flags:
            accepted_sink.add_many(row_number(record), accepted_flags)


def _audit_domain_scale_checks(records: list[RowRecord], sink: IssueSink) -> None:
    for record in records:
        analysis = analyze_record(record)
        codes = list(analysis.get("issue_codes") or [])
        if codes:
            sink.add_many(row_number(record), codes)


def _build_detail_rows(
    records: list[RowRecord],
    detail_map: dict[int, list[str]],
) -> list[dict[str, Any]]:
    row_index = {row_number(r): r for r in records}
    return [
        _detail_row(row_index[rn], sorted(set(codes)))
        for rn, codes in sorted(detail_map.items())
        if rn in row_index
    ]


def _scan_partial_metric_dispositions(
    records: list[RowRecord],
) -> tuple[Counter[str], Counter[str]]:
    """Independent row scan of partial-metric issue and accepted codes."""
    issue_counts: Counter[str] = Counter()
    accepted_counts: Counter[str] = Counter()
    complete_student_contexts = student_contexts_with_complete_metric(records)
    for record in records:
        status = baseline_target_status(record)
        if status not in {"baseline_only", "target_only"}:
            continue
        issue_code, accepted_code = _partial_metric_disposition(
            record,
            status=status,
            complete_student_contexts=complete_student_contexts,
        )
        if issue_code:
            issue_counts[issue_code] += 1
        elif accepted_code:
            accepted_counts[accepted_code] += 1
        else:
            issue_counts["__missing_partial_disposition__"] += 1
    return issue_counts, accepted_counts


def _scan_partial_metric_dispositions_for_status(
    records: list[RowRecord],
    *,
    status: str,
) -> Counter[str]:
    counts: Counter[str] = Counter()
    complete_student_contexts = student_contexts_with_complete_metric(records)
    for record in records:
        if baseline_target_status(record) != status:
            continue
        issue_code, accepted_code = _partial_metric_disposition(
            record,
            status=status,
            complete_student_contexts=complete_student_contexts,
        )
        counts[issue_code or accepted_code or "__missing_partial_disposition__"] += 1
    return counts


def reconcile_audit_counts(
    records: list[RowRecord],
    results: dict[str, Any],
) -> dict[str, Any]:
    errors: list[str] = []
    rows = len(records)
    distribution = results.get("baseline_target_distribution") or {}
    issue_counts = results.get("issue_code_counts") or {}
    accepted_counts = results.get("accepted_exception_counts") or {}
    detail_rows = results.get("detail_rows") or []
    accepted_detail_rows = results.get("accepted_detail_rows") or []

    if sum(int(distribution.get(key, 0) or 0) for key in distribution) != rows:
        errors.append(
            f"baseline_target_distribution sums to {sum(distribution.values())}, expected {rows}"
        )

    scan_issue, scan_accepted = _scan_partial_metric_dispositions(records)
    for code in PARTIAL_METRIC_ISSUE_CODES:
        if scan_issue.get(code, 0) != issue_counts.get(code, 0):
            errors.append(
                f"partial issue {code}: scan={scan_issue.get(code, 0)} "
                f"sink={issue_counts.get(code, 0)}"
            )
    for code in ACCEPTED_ISSUE_CODES:
        if scan_accepted.get(code, 0) != accepted_counts.get(code, 0):
            errors.append(
                f"accepted {code}: scan={scan_accepted.get(code, 0)} "
                f"sink={accepted_counts.get(code, 0)}"
            )
    if scan_issue.get("__missing_partial_disposition__", 0):
        errors.append(
            "partial-metric rows missing disposition: "
            f"{scan_issue['__missing_partial_disposition__']}"
        )

    raw_baseline = int(distribution.get("baseline_without_target", 0) or 0)
    baseline_scan = _scan_partial_metric_dispositions_for_status(
        records,
        status="baseline_only",
    )
    if sum(baseline_scan.values()) != raw_baseline:
        errors.append(
            f"baseline_only rows: scan={sum(baseline_scan.values())} "
            f"distribution={raw_baseline}"
        )

    raw_target = int(distribution.get("target_without_baseline", 0) or 0)
    target_scan = _scan_partial_metric_dispositions_for_status(
        records,
        status="target_only",
    )
    if sum(target_scan.values()) != raw_target:
        errors.append(
            f"target_only rows: scan={sum(target_scan.values())} "
            f"distribution={raw_target}"
        )

    supplemental_excluded = accepted_baseline_supplemental_exclusions(accepted_counts)
    excl = baseline_without_target_excl_supplemental(distribution, accepted_counts)
    flagged_baseline = int(issue_counts.get("baseline_without_target", 0) or 0)
    if excl != flagged_baseline:
        errors.append(
            f"baseline excl supplemental={excl} != flagged={flagged_baseline}"
        )
    if excl + supplemental_excluded != raw_baseline:
        errors.append(
            f"baseline partition: excl={excl} + supplemental_excluded="
            f"{supplemental_excluded} != raw={raw_baseline}"
        )

    if sum(int(accepted_counts.get(code, 0) or 0) for code in ACCEPTED_ISSUE_CODES) != len(
        accepted_detail_rows
    ):
        errors.append(
            "accepted code count sum does not match accepted_detail_rows length"
        )

    partial_counts, data_quality_counts = partition_issue_code_counts(issue_counts)
    if sum(partial_counts.values()) + sum(data_quality_counts.values()) != sum(
        issue_counts.values()
    ):
        errors.append("issue_code_counts partition does not cover all issue codes")

    partial_row_count = count_rows_with_issue_codes(detail_rows, PARTIAL_METRIC_ISSUE_CODES)
    if partial_row_count != count_rows_with_issue_codes(
        detail_rows,
        frozenset(partial_counts),
    ):
        errors.append("partial_metric_issue_counts disagree with detail_rows")

    metrics = year_audit_metrics(results)
    if metrics["baseline_without_target_excl_supplemental"] != excl:
        errors.append("year_audit_metrics baseline_without_target_excl_supplemental mismatch")
    if metrics["partial_goal_metrics_flagged"] != flagged_baseline:
        errors.append("year_audit_metrics partial_goal_metrics_flagged mismatch")
    if metrics["accepted_exceptions_total"] != len(accepted_detail_rows):
        errors.append("year_audit_metrics accepted_exceptions_total mismatch")
    if metrics["rows_evaluated"] != rows:
        errors.append("year_audit_metrics rows_evaluated mismatch")

    overlap = {
        int(row["row_number"])
        for row in detail_rows
        if PARTIAL_METRIC_ISSUE_CODES
        & {code.strip() for code in (row.get("issue_codes") or "").split(";")}
    } & {int(row["row_number"]) for row in accepted_detail_rows}
    if overlap:
        errors.append(f"rows appear in both flagged and accepted sinks: {sorted(overlap)}")

    return {
        "ok": not errors,
        "error_count": len(errors),
        "errors": errors,
        "baseline_without_target_raw": raw_baseline,
        "baseline_without_target_excl_supplemental": excl,
        "baseline_supplemental_excluded": supplemental_excluded,
        "baseline_flagged_primary": flagged_baseline,
    }


def run_audit(
    records: list[RowRecord],
    *,
    school_year: str | None = None,
    excluded_student_ids: set[str] | None = None,
    goal_progress_workbook: Path | None = None,
    final_goal_achievement_students: set[str] | None = None,
) -> dict[str, Any]:
    del final_goal_achievement_students  # ponytail: legacy kwarg; EWGSPE scope replaced GA drilldown
    resolved_excluded = excluded_student_ids
    if resolved_excluded is None and school_year is not None:
        resolved_excluded = resolve_ewgspe_only_student_ids(
            school_year,
            goal_progress_workbook=goal_progress_workbook,
        )
    excluded = resolved_excluded or set()
    records, ewgspe_rows_excluded = filter_records_excluding_students(records, excluded)
    summary: Counter[str] = Counter()
    issue_codes: Counter[str] = Counter()
    detail_map: dict[int, list[str]] = defaultdict(list)
    accepted_codes: Counter[str] = Counter()
    accepted_detail_map: dict[int, list[str]] = defaultdict(list)
    sink = IssueSink(issue_codes=issue_codes, detail_map=detail_map)
    accepted_sink = AcceptedSink(
        accepted_codes=accepted_codes,
        detail_map=accepted_detail_map,
    )

    key_info = audit_composite_keys(records)
    client_info = audit_student_client_ids(records)
    complete_student_contexts = student_contexts_with_complete_metric(records)
    bt = audit_baseline_target(records)

    summary["rows"] = len(records)
    summary["ewgspe_only_students_excluded"] = len(excluded)
    summary["rows_excluded_ewgspe_students"] = ewgspe_rows_excluded
    summary["exited_rows_excluded"] = 0
    summary["exited_rows_included_with_final_goal_achievement"] = 0
    if school_year is not None:
        summary["enrollment_filter"] = "all_students_except_ewgspe_only"
    else:
        summary["enrollment_filter"] = "all_rows"
    summary["duplicate_composite_keys"] = key_info["duplicate_count"]
    summary["student_client_id_conflicts"] = client_info["student_ids_with_multiple_client_ids"]
    for key, value in bt.items():
        summary[key] = value

    _audit_case_manager_blanks(records, sink)
    _mark_duplicate_composite_keys(key_info, sink)
    _mark_student_client_id_mismatches(
        records,
        conflict_sid_set=set(client_info.get("conflicting_student_ids") or []),
        sink=sink,
    )
    _audit_baseline_target_flags(
        records,
        complete_student_contexts=complete_student_contexts,
        sink=sink,
        accepted_sink=accepted_sink,
    )
    _audit_domain_scale_checks(records, sink)
    student_completeness_codes: Counter[str] = Counter()
    student_completeness_map: dict[int, list[str]] = defaultdict(list)
    if school_year is not None:
        summary["student_completeness_flagged"] = _audit_student_abc_completeness(
            records,
            issue_codes=student_completeness_codes,
            detail_map=student_completeness_map,
        )
    else:
        summary["student_completeness_flagged"] = 0

    detail_rows = _build_detail_rows(records, detail_map)
    accepted_detail_rows = _build_detail_rows(records, accepted_detail_map)
    student_completeness_rows = _build_detail_rows(records, student_completeness_map)
    summary["accepted_exception_rows"] = len(accepted_detail_rows)

    results = {
        "composite_key_audit": key_info,
        "student_client_audit": client_info,
        "baseline_target_distribution": bt,
        "summary": dict(summary),
        "issue_code_counts": dict(issue_codes),
        "accepted_exception_counts": dict(accepted_codes),
        "detail_rows": detail_rows,
        "accepted_detail_rows": accepted_detail_rows,
        "student_completeness_rows": student_completeness_rows,
        "student_completeness_issue_counts": dict(student_completeness_codes),
        "progress_rollup": progress_rollup_from_records(records),
    }
    results["reconciliation"] = reconcile_audit_counts(records, results)
    return results


def _detail_row(record: RowRecord, codes: list[str]) -> dict[str, Any]:
    return {
        "row_number": record.get("row_number"),
        "student_id": clean_value(_text(record, "Student ID")),
        "school": clean_value(_text(record, "School")),
        "goal": clean_value(_text(record, "Goal")),
        "metric": clean_value(_text(record, "Metric")),
        "baseline": clean_value(_text(record, "Baseline")),
        "target": clean_value(_text(record, "Target")),
        "latest_progress": latest_metric_value(record),
        "issue_codes": ";".join(codes),
    }


def write_audit_flags_csv(
    path: Path,
    rows: list[dict[str, Any]],
    *,
    include_school_year: bool = False,
) -> None:
    fieldnames = [
        "row_number",
        "student_id",
        "school",
        "goal",
        "metric",
        "baseline",
        "target",
        "latest_progress",
        "issue_codes",
    ]
    if include_school_year:
        fieldnames.insert(0, "school_year")
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        for row in rows:
            writer.writerow({k: row.get(k) for k in fieldnames})


def partition_rows_matching_codes(
    rows: list[dict[str, Any]],
    codes: frozenset[str],
) -> list[dict[str, Any]]:
    partitioned: list[dict[str, Any]] = []
    for row in rows:
        row_codes = {
            code.strip()
            for code in (row.get("issue_codes") or "").split(";")
            if code.strip()
        }
        matched = row_codes & codes
        if matched:
            partitioned.append({**row, "issue_codes": ";".join(sorted(matched))})
    return partitioned


def partition_flagged_detail_rows(
    rows: list[dict[str, Any]],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    partial_rows = partition_rows_matching_codes(rows, PARTIAL_METRIC_ISSUE_CODES)
    data_quality_rows: list[dict[str, Any]] = []
    for row in rows:
        row_codes = {
            code.strip()
            for code in (row.get("issue_codes") or "").split(";")
            if code.strip()
        }
        data_quality_codes = row_codes - PARTIAL_METRIC_ISSUE_CODES
        if data_quality_codes:
            data_quality_rows.append(
                {**row, "issue_codes": ";".join(sorted(data_quality_codes))}
            )
    return partial_rows, data_quality_rows


def partition_accepted_detail_rows(
    rows: list[dict[str, Any]],
) -> dict[str, list[dict[str, Any]]]:
    buckets = {export_key: [] for export_key, _, _ in ACCEPTED_EXPORT_SPECS}
    for row in rows:
        row_codes = {
            code.strip()
            for code in (row.get("issue_codes") or "").split(";")
            if code.strip()
        }
        for export_key, codes, _filename in ACCEPTED_EXPORT_SPECS:
            if row_codes & codes:
                buckets[export_key].append(row)
                break
    return buckets


def write_partitioned_audit_exports(
    output_dir: Path,
    *,
    detail_rows: list[dict[str, Any]],
    accepted_detail_rows: list[dict[str, Any]],
    student_completeness_rows: list[dict[str, Any]] | None = None,
    include_school_year: bool = False,
) -> dict[str, str]:
    output_dir.mkdir(parents=True, exist_ok=True)
    partial_rows, data_quality_rows = partition_flagged_detail_rows(detail_rows)
    accepted_buckets = partition_accepted_detail_rows(accepted_detail_rows)
    paths: dict[str, str] = {
        "audit_flags": str(output_dir / "audit_flags.csv"),
        "partial_goal_metrics_flags": str(output_dir / "partial_goal_metrics_flags.csv"),
        "student_completeness_flags": str(output_dir / "student_completeness_flags.csv"),
    }
    write_audit_flags_csv(
        Path(paths["audit_flags"]),
        data_quality_rows,
        include_school_year=include_school_year,
    )
    write_audit_flags_csv(
        Path(paths["partial_goal_metrics_flags"]),
        partial_rows,
        include_school_year=include_school_year,
    )
    write_audit_flags_csv(
        Path(paths["student_completeness_flags"]),
        student_completeness_rows or [],
        include_school_year=include_school_year,
    )
    for export_key, _codes, filename in ACCEPTED_EXPORT_SPECS:
        path = output_dir / filename
        paths[export_key] = str(path)
        write_audit_flags_csv(
            path,
            accepted_buckets[export_key],
            include_school_year=include_school_year,
        )
    return paths


def _format_issue_counts(counts: dict[str, int]) -> list[str]:
    lines = []
    for code in sorted(counts):
        label = FLAG_DESCRIPTIONS.get(code, code)
        lines.append(f"- **{code}** ({label}): {counts[code]}")
    return lines


def partition_issue_code_counts(
    counts: dict[str, int],
) -> tuple[dict[str, int], dict[str, int]]:
    partial_metric: dict[str, int] = {}
    data_quality: dict[str, int] = {}
    for code, count in counts.items():
        if code in PARTIAL_METRIC_ISSUE_CODES:
            partial_metric[code] = count
        else:
            data_quality[code] = count
    return partial_metric, data_quality


def count_rows_with_issue_codes(
    detail_rows: list[dict[str, Any]],
    codes: frozenset[str],
) -> int:
    matched = 0
    for row in detail_rows:
        row_codes = {code.strip() for code in (row.get("issue_codes") or "").split(";")}
        row_codes.discard("")
        if row_codes & codes:
            matched += 1
    return matched


def accepted_baseline_supplemental_exclusions(accepted_counts: dict[str, int]) -> int:
    return sum(
        int(accepted_counts.get(code, 0) or 0) for code in BASELINE_SUPPLEMENTAL_ACCEPTED_CODES
    )


def baseline_without_target_excl_supplemental(
    distribution: dict[str, int],
    accepted_counts: dict[str, int],
) -> int:
    raw = int(distribution.get("baseline_without_target", 0) or 0)
    return raw - accepted_baseline_supplemental_exclusions(accepted_counts)


def _format_baseline_without_target_breakdown(
    *,
    distribution: dict[str, int],
    partial_counts: dict[str, int],
    accepted_counts: dict[str, int],
) -> str | None:
    raw = int(distribution.get("baseline_without_target", 0) or 0)
    if not raw:
        return None
    excluded = accepted_baseline_supplemental_exclusions(accepted_counts)
    total = raw - excluded
    flagged = int(partial_counts.get("baseline_without_target", 0) or 0)
    return (
        f"- Baseline without target (excl. non-goal metrics): **{total}** "
        f"({flagged} flagged; {excluded} non-goal excluded)"
    )


def _format_partial_metric_breakdown(
    *,
    distribution_key: str,
    distribution: dict[str, int],
    flagged_counts: dict[str, int],
    accepted_counts: dict[str, int],
    flagged_code: str,
    accepted_code: str,
    label: str,
) -> str | None:
    total = distribution.get(distribution_key, 0)
    if not total:
        return None
    flagged = flagged_counts.get(flagged_code, 0)
    accepted = accepted_counts.get(accepted_code, 0)
    return f"- {label}: **{total}** ({flagged} flagged, {accepted} accepted)"


def write_summary_markdown(path: Path, payload: dict[str, Any]) -> None:
    counts = payload.get("issue_code_counts") or {}
    partial_counts = payload.get("partial_metric_issue_counts")
    data_quality_counts = payload.get("data_quality_issue_counts")
    if partial_counts is None or data_quality_counts is None:
        partial_counts, data_quality_counts = partition_issue_code_counts(counts)
    accepted_counts = payload.get("accepted_exception_counts") or {}
    summary = payload.get("summary") or {}
    distribution = payload.get("baseline_target_distribution") or {}
    detail_rows = payload.get("detail_rows") or []
    accepted_row_count = payload.get("accepted_detail_row_count")
    if accepted_row_count is None:
        accepted_row_count = len(payload.get("accepted_detail_rows") or [])
    partial_row_count = summary.get("partial_metric_flagged_rows")
    if partial_row_count is None:
        partial_row_count = count_rows_with_issue_codes(
            detail_rows,
            PARTIAL_METRIC_ISSUE_CODES,
        )
    data_quality_row_count = summary.get("data_quality_flagged_rows")
    if data_quality_row_count is None:
        data_quality_row_count = count_rows_with_issue_codes(
            detail_rows,
            frozenset(data_quality_counts),
        )
    lines = [
        "# Student Metrics Audit Summary",
        "",
        f"- **Workbook:** `{payload.get('workbook', '')}`",
        f"- **Sheet:** {payload.get('sheet', DEFAULT_SHEET)}",
        f"- **Run at:** {payload.get('run_at', '')}",
        "",
        "## Counts",
        "",
        f"- Rows evaluated: **{summary.get('rows', 0)}**",
        (
            "- EWGSPE-only students excluded: "
            f"**{summary.get('ewgspe_only_students_excluded', 0)}** "
            f"({summary.get('rows_excluded_ewgspe_students', 0)} rows)"
        ),
        f"- Enrollment scope: **{summary.get('enrollment_filter', 'all_rows')}**",
        (
            "- Student ABC completeness flagged: "
            f"**{summary.get('student_completeness_flagged', 0)}**"
        ),
        f"- Rows flagged: **{len(detail_rows)}**",
        f"- Partial goal metrics (flagged): **{partial_row_count}**",
        f"- Data quality issues: **{data_quality_row_count}**",
        f"- Accepted exceptions: **{accepted_row_count}**",
        f"- Duplicate composite keys: **{summary.get('duplicate_composite_keys', 0)}**",
        f"- Student/client ID conflicts: **{summary.get('student_client_id_conflicts', 0)}**",
        "",
    ]
    reconciliation = payload.get("reconciliation") or {}
    if reconciliation:
        status = "passed" if reconciliation.get("ok") else "FAILED"
        lines.append(f"- Reconciliation: **{status}**")
        lines.append("")
    for breakdown in (
        _format_baseline_without_target_breakdown(
            distribution=distribution,
            partial_counts=partial_counts,
            accepted_counts=accepted_counts,
        ),
        _format_partial_metric_breakdown(
            distribution_key="target_without_baseline",
            distribution=distribution,
            flagged_counts=partial_counts,
            accepted_counts=accepted_counts,
            flagged_code="target_without_baseline",
            accepted_code="target_without_baseline_supplemental_metric",
            label="Target without baseline",
        ),
    ):
        if breakdown:
            lines.append(breakdown)
    if any(
        line.startswith("- Baseline without target (excl. non-goal metrics):")
        or line.startswith("- Target without baseline:")
        for line in lines
    ):
        lines.append("")
    if accepted_counts:
        lines.extend(["## Accepted exceptions", ""])
        lines.append(
            "Non-goal metrics: baseline-only rows for students who already have another "
            "complete goal-metric row in the same school year."
        )
        lines.append("")
        lines.extend(_format_issue_counts(accepted_counts))
        lines.append("")
    if partial_counts:
        lines.extend(["## Partial goal metrics", ""])
        lines.append(
            "Baseline-only or target-only rows where the student has no other complete "
            "goal-metric row in the school year. Row-level detail: "
            "`partial_goal_metrics_flags.csv`."
        )
        lines.append("")
        lines.extend(_format_issue_counts(partial_counts))
        lines.append("")
    lines.extend(["## Data quality issues", ""])
    if data_quality_counts:
        lines.append("Row-level detail: `audit_flags.csv`.")
        lines.append("")
        lines.extend(_format_issue_counts(data_quality_counts))
    else:
        lines.append("No data quality issues found.")
    lines.append("")
    path.write_text("\n".join(lines), encoding="utf-8")


def _pct(part: int, whole: int) -> str:
    if whole <= 0:
        return "0%"
    return f"{round(100 * part / whole)}%"


def year_audit_metrics(results: dict[str, Any]) -> dict[str, Any]:
    rows = int(results.get("summary", {}).get("rows", 0) or 0)
    distribution = results.get("baseline_target_distribution") or {}
    issue_counts = results.get("issue_code_counts") or {}
    accepted_counts = results.get("accepted_exception_counts") or {}
    detail_rows = results.get("detail_rows") or []
    accepted_rows = len(results.get("accepted_detail_rows") or [])
    partial_counts, data_quality_counts = partition_issue_code_counts(issue_counts)
    baseline_and_target = int(distribution.get("baseline_and_target_present", 0) or 0)
    baseline_excl_supplemental = baseline_without_target_excl_supplemental(
        distribution,
        accepted_counts,
    )
    domain_tracking = int(
        accepted_counts.get("baseline_without_target_domain_tracking", 0)
        + accepted_counts.get("target_without_baseline_domain_tracking", 0)
    )
    supplemental = int(
        accepted_counts.get("baseline_without_target_supplemental_metric", 0)
        + accepted_counts.get("target_without_baseline_supplemental_metric", 0)
    )
    no_abc_goals = int(
        accepted_counts.get("baseline_without_target_no_goal_context", 0)
        + accepted_counts.get("target_without_baseline_no_goal_context", 0)
    )
    return {
        "rows_evaluated": rows,
        "baseline_and_target_present": baseline_and_target,
        "baseline_without_target_excl_supplemental": baseline_excl_supplemental,
        "partial_goal_metrics_flagged": int(
            issue_counts.get("baseline_without_target", 0) or 0
        ),
        "accepted_domain_tracking": domain_tracking,
        "accepted_supplemental": supplemental,
        "accepted_no_abc_goals": no_abc_goals,
        "accepted_exceptions_total": accepted_rows,
        "rows_flagged": len(detail_rows),
        "partial_metric_rows_flagged": count_rows_with_issue_codes(
            detail_rows,
            PARTIAL_METRIC_ISSUE_CODES,
        ),
        "data_quality_rows_flagged": count_rows_with_issue_codes(
            detail_rows,
            frozenset(data_quality_counts),
        ),
        "direction_mismatches": int(issue_counts.get("baseline_target_direction_mismatch", 0) or 0),
        "scale_mismatches": int(issue_counts.get("baseline_scale_mismatch", 0) or 0)
        + int(issue_counts.get("target_scale_mismatch", 0) or 0),
        "duplicate_composite_keys": int(results.get("summary", {}).get("duplicate_composite_keys", 0) or 0),
        "student_client_id_conflicts": int(
            results.get("summary", {}).get("student_client_id_conflicts", 0) or 0
        ),
    }


def build_dataset_manifest(
    *,
    export_paths: dict[str, str],
    detail_rows: list[dict[str, Any]],
    accepted_detail_rows: list[dict[str, Any]],
    student_completeness_rows: list[dict[str, Any]] | None = None,
) -> dict[str, dict[str, Any]]:
    partial_rows, data_quality_rows = partition_flagged_detail_rows(detail_rows)
    accepted_buckets = partition_accepted_detail_rows(accepted_detail_rows)
    manifest: dict[str, dict[str, Any]] = {
        "audit_flags": {
            "path": Path(export_paths["audit_flags"]).name,
            "row_count": len(data_quality_rows),
        },
        "partial_goal_metrics_flags": {
            "path": Path(export_paths["partial_goal_metrics_flags"]).name,
            "row_count": len(partial_rows),
        },
        "student_completeness_flags": {
            "path": Path(export_paths["student_completeness_flags"]).name,
            "row_count": len(student_completeness_rows or []),
        },
    }
    for export_key, _codes, filename in ACCEPTED_EXPORT_SPECS:
        manifest[export_key] = {
            "path": filename,
            "row_count": len(accepted_buckets[export_key]),
        }
    return manifest


def build_cumulative_payload(
    year_runs: list[tuple[str, dict[str, Any], str]],
) -> dict[str, Any]:
    school_years = [school_year for school_year, _, _ in year_runs]
    per_year: dict[str, dict[str, Any]] = {}
    comparison: dict[str, dict[str, Any]] = {}
    for school_year, results, workbook in year_runs:
        metrics = year_audit_metrics(results)
        per_year[school_year] = {
            "workbook": workbook,
            "sheet": results.get("sheet", DEFAULT_SHEET),
            "metrics": metrics,
            "reconciliation": results.get("reconciliation") or {},
            "issue_code_counts": results.get("issue_code_counts") or {},
            "accepted_exception_counts": results.get("accepted_exception_counts") or {},
            "baseline_target_distribution": results.get("baseline_target_distribution") or {},
            "summary": results.get("summary") or {},
        }
        for key, value in metrics.items():
            comparison.setdefault(key, {})[school_year] = value

    return {
        "cumulative": True,
        "run_at": datetime.now(timezone.utc).isoformat(),
        "school_years": school_years,
        "reconciliation_ok": all(
            (per_year[school_year].get("reconciliation") or {}).get("ok", False)
            for school_year in school_years
        ),
        "years": per_year,
        "comparison": comparison,
    }


def _format_cumulative_cell(metric_key: str, value: Any, *, rows: int) -> str:
    if metric_key in {
        "baseline_and_target_present",
        "baseline_without_target_excl_supplemental",
    } and rows > 0:
        return f"{value} ({_pct(int(value), rows)})"
    return str(value)


def write_cumulative_summary_markdown(path: Path, payload: dict[str, Any]) -> None:
    school_years = payload.get("school_years") or []
    comparison = payload.get("comparison") or {}
    rows_by_year = {
        school_year: int((payload.get("years", {}).get(school_year, {}).get("metrics") or {}).get("rows_evaluated", 0))
        for school_year in school_years
    }
    metric_labels = {
        "rows_evaluated": "Rows evaluated",
        "baseline_and_target_present": "Baseline + target present",
        "baseline_without_target_excl_supplemental": (
            "Baseline without target (excl. non-goal metrics)"
        ),
        "partial_goal_metrics_flagged": (
            "Baseline without target flagged (no other complete goal in SY)"
        ),
        "accepted_domain_tracking": "Accepted: domain tracking metrics",
        "accepted_supplemental": "Accepted: supplemental metrics",
        "accepted_no_abc_goals": "Accepted: no ABC goals set yet",
        "accepted_exceptions_total": "Accepted exception rows (total)",
        "rows_flagged": "Rows flagged (all issues)",
        "direction_mismatches": "Direction mismatches",
        "scale_mismatches": "Scale mismatches",
    }
    lines = [
        "# Student Metrics Audit (cumulative)",
        "",
        "Rules: include all students except those whose Goal Achievement is only "
        f'"{GOAL_ACHIEVEMENT_EWGSPE}" (from Goal Progress Detail when available); '
        "exclude baseline-only rows when the student has another complete "
        "goal-metric row in the same school year (non-goal metrics); flag baseline-only "
        "or target-only rows when no complete goal exists for that student in the SY; "
        "student ABC completeness (3 baselines + 1 target, attendance-only for Pocono/SE Delco) "
        "in `student_completeness_flags.csv`.",
        "",
        f"- **Run at:** {payload.get('run_at', '')}",
        f"- **School years:** {', '.join(school_years)}",
        "",
        "| Metric | " + " | ".join(school_years) + " |",
        "|---|" + "|".join(["---:"] * len(school_years)) + "|",
    ]
    for key in CUMULATIVE_METRIC_ORDER:
        if key not in comparison:
            continue
        label = metric_labels.get(key, key)
        cells = [
            _format_cumulative_cell(
                key,
                comparison[key].get(school_year, 0),
                rows=rows_by_year.get(school_year, 0),
            )
            for school_year in school_years
        ]
        lines.append(f"| {label} | " + " | ".join(cells) + " |")
    reconciliation_ok = payload.get("reconciliation_ok")
    if reconciliation_ok is not None:
        lines.extend(
            [
                "",
                f"- **Count reconciliation:** {'passed' if reconciliation_ok else 'FAILED'}",
            ]
        )
    lines.append("")
    path.write_text("\n".join(lines), encoding="utf-8")


def export_cumulative_results(
    year_runs: list[tuple[str, dict[str, Any], str]],
    output_dir: Path,
) -> dict[str, str]:
    output_dir.mkdir(parents=True, exist_ok=True)
    payload = build_cumulative_payload(year_runs)
    detail_rows: list[dict[str, Any]] = []
    accepted_detail_rows: list[dict[str, Any]] = []
    student_completeness_rows: list[dict[str, Any]] = []
    for school_year, results, _workbook in year_runs:
        for row in results.get("detail_rows") or []:
            detail_rows.append({"school_year": school_year, **row})
        for row in results.get("accepted_detail_rows") or []:
            accepted_detail_rows.append({"school_year": school_year, **row})
        for row in results.get("student_completeness_rows") or []:
            student_completeness_rows.append({"school_year": school_year, **row})
    export_paths = write_partitioned_audit_exports(
        output_dir,
        detail_rows=detail_rows,
        accepted_detail_rows=accepted_detail_rows,
        student_completeness_rows=student_completeness_rows,
        include_school_year=True,
    )
    paths = {
        **export_paths,
        "audit_summary_json": str(output_dir / "audit_summary.json"),
        "audit_summary_md": str(output_dir / "audit_summary.md"),
    }
    payload["datasets"] = build_dataset_manifest(
        export_paths=export_paths,
        detail_rows=detail_rows,
        accepted_detail_rows=accepted_detail_rows,
        student_completeness_rows=student_completeness_rows,
    )
    Path(paths["audit_summary_json"]).write_text(
        json.dumps(payload, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    write_cumulative_summary_markdown(Path(paths["audit_summary_md"]), payload)
    return paths


def results_to_json_payload(
    results: dict[str, Any],
    *,
    workbook: str,
    sheet: str,
    school_year: str | None = None,
) -> dict[str, Any]:
    issue_counts = results.get("issue_code_counts", {})
    partial_counts, data_quality_counts = partition_issue_code_counts(issue_counts)
    detail_rows = results.get("detail_rows") or []
    summary = dict(results.get("summary") or {})
    summary["partial_metric_flagged_rows"] = count_rows_with_issue_codes(
        detail_rows,
        PARTIAL_METRIC_ISSUE_CODES,
    )
    summary["data_quality_flagged_rows"] = count_rows_with_issue_codes(
        detail_rows,
        frozenset(data_quality_counts),
    )
    payload = {
        "workbook": workbook,
        "sheet": sheet,
        "run_at": datetime.now(timezone.utc).isoformat(),
        "summary": summary,
        "issue_code_counts": issue_counts,
        "partial_metric_issue_counts": partial_counts,
        "data_quality_issue_counts": data_quality_counts,
        "accepted_exception_counts": results.get("accepted_exception_counts", {}),
        "baseline_target_distribution": results.get("baseline_target_distribution", {}),
        "composite_key_audit": results.get("composite_key_audit", {}),
        "student_client_audit": results.get("student_client_audit", {}),
        "detail_row_count": len(detail_rows),
        "accepted_detail_row_count": len(results.get("accepted_detail_rows") or []),
        "student_completeness_row_count": len(results.get("student_completeness_rows") or []),
        "student_completeness_issue_counts": results.get("student_completeness_issue_counts", {}),
        "reconciliation": results.get("reconciliation") or {},
    }
    if school_year:
        payload["school_year"] = school_year
    return payload


def export_results(
    results: dict[str, Any],
    output_dir: Path,
    *,
    workbook: str,
    sheet: str,
    school_year: str | None = None,
) -> dict[str, str]:
    output_dir.mkdir(parents=True, exist_ok=True)
    detail_rows = results.get("detail_rows") or []
    accepted_detail_rows = results.get("accepted_detail_rows") or []
    student_completeness_rows = results.get("student_completeness_rows") or []
    export_paths = write_partitioned_audit_exports(
        output_dir,
        detail_rows=detail_rows,
        accepted_detail_rows=accepted_detail_rows,
        student_completeness_rows=student_completeness_rows,
    )
    payload = results_to_json_payload(
        results,
        workbook=workbook,
        sheet=sheet,
        school_year=school_year,
    )
    payload["datasets"] = build_dataset_manifest(
        export_paths=export_paths,
        detail_rows=detail_rows,
        accepted_detail_rows=accepted_detail_rows,
        student_completeness_rows=student_completeness_rows,
    )
    paths = {
        **export_paths,
        "audit_summary_json": str(output_dir / "audit_summary.json"),
        "audit_summary_md": str(output_dir / "audit_summary.md"),
    }
    Path(paths["audit_summary_json"]).write_text(
        json.dumps(payload, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    write_summary_markdown(
        Path(paths["audit_summary_md"]),
        {
            **payload,
            "detail_rows": detail_rows,
        },
    )
    return paths


def evaluate_workbook(workbook_path: Path, config: AuditConfig) -> dict[str, Any]:
    records = load_rows(workbook_path, sheet_name=config.sheet_name)
    results = run_audit(
        records,
        school_year=config.school_year,
        excluded_student_ids=config.excluded_student_ids,
        goal_progress_workbook=config.goal_progress_workbook,
    )
    results["workbook"] = str(workbook_path.resolve())
    results["sheet"] = config.sheet_name
    results["config"] = asdict(config)
    return results


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Flag invalid metric/scale combinations in student metrics summary exports.",
    )
    parser.add_argument(
        "--year",
        action="append",
        nargs="+",
        dest="years",
        metavar="SY",
        help="School year(s) to audit. Use 'all' for a cumulative cross-year report.",
    )
    parser.add_argument(
        "--workbook",
        action="append",
        type=Path,
        help="Path to student metrics .xlsx (repeatable). When omitted, resolves via cisiphyus.",
    )
    parser.add_argument(
        "--force-fetch",
        action="store_true",
        help="Always re-export from CISDM even when the local workbook is fresh.",
    )
    parser.add_argument(
        "--fetch-destination",
        type=Path,
        help="Where to save the CISDM export (default: artifacts/audit/).",
    )
    parser.add_argument(
        "--sheet",
        default=DEFAULT_SHEET,
        help=f"Worksheet name (default {DEFAULT_SHEET}).",
    )
    parser.add_argument(
        "--strict",
        action="store_true",
        help="Exit 1 on duplicate keys, ID conflicts, or count reconciliation failures.",
    )
    parser.add_argument(
        "--json",
        action="store_true",
        help="Print JSON summary to stdout.",
    )
    return parser


def _dedupe_years(years: list[str]) -> list[str]:
    seen: set[str] = set()
    unique: list[str] = []
    for year in years:
        if year not in seen:
            seen.add(year)
            unique.append(year)
    return unique


def _flatten_year_tokens(groups: list[list[str]] | None) -> list[str]:
    if not groups:
        return []
    tokens: list[str] = []
    for group in groups:
        tokens.extend(group)
    return tokens


def _resolve_audit_plan(
    args: argparse.Namespace,
) -> tuple[Literal["cumulative", "per_year", "workbook"], list[tuple[Path, str | None]]]:
    if args.workbook and args.years:
        raise ValueError("--workbook and --year are mutually exclusive")

    if args.workbook:
        return "workbook", [(path.resolve(), None) for path in args.workbook]

    catalog = discover_audit_school_years()
    tokens = _flatten_year_tokens(args.years)
    cumulative = any(token.strip().lower() == "all" for token in tokens)
    if cumulative and len(tokens) > 1:
        raise ValueError("use --year all alone for cumulative audit")

    years_to_run: list[str] = []
    if cumulative:
        years_to_run = catalog
        mode: Literal["cumulative", "per_year"] = "cumulative"
    elif tokens:
        years_to_run = [normalize_school_year(token, catalog) for token in tokens]
        mode = "per_year"
    else:
        school_year = current_school_year()
        if not school_year:
            raise ValueError("no --year provided and no default school year in config")
        years_to_run = [school_year]
        mode = "per_year"

    jobs: list[tuple[Path, str | None]] = []
    for school_year in _dedupe_years(years_to_run):
        workbook = resolve_workbook_for_school_year(
            school_year,
            force_fetch=args.force_fetch,
            fetch_destination=args.fetch_destination,
        )
        jobs.append((workbook, school_year))
    return mode, jobs


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    config = AuditConfig(sheet_name=args.sheet)

    try:
        mode, jobs = _resolve_audit_plan(args)
    except ValueError as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 2
    except (FileNotFoundError, OSError, RuntimeError) as exc:
        print(f"error: CISDM fetch failed: {exc}", file=sys.stderr)
        return 1

    if not jobs:
        print("error: no workbook available", file=sys.stderr)
        return 1

    base_out = default_output_dir()
    year_runs: list[tuple[str, dict[str, Any], str]] = []
    summaries: list[tuple[dict[str, Any], str | None]] = []

    for workbook, school_year in jobs:
        if not workbook.is_file():
            print(f"error: workbook not found: {workbook}", file=sys.stderr)
            return 1
        year_config = AuditConfig(
            sheet_name=config.sheet_name,
            include_ok_rows=config.include_ok_rows,
            school_year=school_year,
        )
        results = evaluate_workbook(workbook, year_config)
        summaries.append((results, school_year))
        if mode == "cumulative":
            year_runs.append((school_year or "", results, str(workbook)))
            continue
        out_dir = base_out / school_year if school_year else base_out
        export_results(
            results,
            out_dir,
            workbook=str(workbook),
            sheet=config.sheet_name,
            school_year=school_year,
        )
        label = school_year or str(workbook)
        print(f"School year: {label}")
        print(f"Workbook: {workbook}")
        print(f"Output dir: {out_dir}")
        print(f"  rows: {results['summary'].get('rows', 0)}")
        print(f"  flagged_rows: {len(results.get('detail_rows') or [])}")
        reconciliation = results.get("reconciliation") or {}
        if reconciliation.get("ok"):
            print("  reconciliation: ok")
        else:
            print(f"  reconciliation: FAILED ({reconciliation.get('error_count', 0)} errors)")
            for message in reconciliation.get("errors") or []:
                print(f"    - {message}", file=sys.stderr)

    if mode == "cumulative":
        paths = export_cumulative_results(year_runs, base_out)
        print(f"Cumulative audit: {len(year_runs)} school years")
        print(f"Output dir: {base_out}")
        for school_year, results, workbook in year_runs:
            reconciliation = results.get("reconciliation") or {}
            recon = "ok" if reconciliation.get("ok") else "FAILED"
            print(
                f"  {school_year}: rows={results['summary'].get('rows', 0)} "
                f"flagged={len(results.get('detail_rows') or [])} "
                f"reconciliation={recon} workbook={workbook}"
            )
        print(f"  summary: {paths['audit_summary_md']}")

    if args.strict:
        for results, _school_year in summaries:
            if results["composite_key_audit"]["duplicate_count"] > 0:
                return 1
            if results["student_client_audit"]["student_ids_with_multiple_client_ids"] > 0:
                return 1
            reconciliation = results.get("reconciliation") or {}
            if not reconciliation.get("ok", False):
                return 1

    if args.json:
        if mode == "cumulative":
            print(json.dumps(build_cumulative_payload(year_runs), indent=2, sort_keys=True))
        else:
            payloads = [
                results_to_json_payload(
                    results,
                    workbook=results["workbook"],
                    sheet=results["sheet"],
                    school_year=school_year,
                )
                for results, school_year in summaries
            ]
            print(json.dumps(payloads if len(payloads) > 1 else payloads[0], indent=2, sort_keys=True))

    return 0


# --- goal achievement audit (GAR) ---

GAR_HEADER_COUNT = 36
GAR_REQUIRED_HEADERS = frozenset(
    {
        "Home School",
        "Metric",
        "Baseline",
        "Target",
        "Goal Achievement",
        "Achieved Value",
    }
)
GAR_METRIC_DIRECTIONS: dict[str, str] = {
    "Attendance Rate (%)": "higher",
    "Attendance Rate (days absent)": "lower",
    "Credits Needed/Credit Completion": "higher",
    "Suspensions": "lower",
    "Tardies": "lower",
    "Disciplinary Referrals": "lower",
    "Conduct": "lower",
    "Other Behavior Incidents": "lower",
    "Drank Alcohol (# of times past 30 days)": "lower",
    "Violent/bullied someone (# of times past 30 days)": "lower",
    "Other (SEL)": "higher",
    "Accept a position of employment": "higher",
    "Create a resume": "higher",
    "Interview with one or more potential employers": "higher",
    "Other High Risk Behavior": "lower",
    "Complete career assessment": "higher",
    "Other College Readiness": "higher",
    "Other Career Readiness": "higher",
    "Apply to one or more colleges/universities": "higher",
    "Accepted to one or more colleges/universities": "higher",
}
GAR_GRADE_SCALE: dict[str, int] = {
    "F": 0,
    "D-": 1,
    "D": 2,
    "D+": 3,
    "C-": 4,
    "C": 5,
    "C+": 6,
    "B-": 7,
    "B": 8,
    "B+": 9,
    "A-": 10,
    "A": 11,
    "A+": 12,
}
GAR_PERCENT_TO_GRADE_THRESHOLD: list[tuple[int, str]] = [
    (100, "A+"),
    (93, "A"),
    (90, "A-"),
    (87, "B+"),
    (83, "B"),
    (80, "B-"),
    (77, "C+"),
    (73, "C"),
    (70, "C-"),
    (67, "D+"),
    (63, "D"),
    (60, "D-"),
    (0, "F"),
]
GAR_ALTERNATE_GRADE_SCALE: dict[str, int] = {
    "BB": 1,
    "BASIC": 2,
    "B": 2,
    "PROFECIENT": 3,
    "PROFICIENT": 3,
    "P": 3,
    "ADVANCED": 4,
    "A": 4,
}
GAR_YES_NO_METRICS = frozenset(
    {
        "Accept a position of employment",
        "Create a resume",
        "Interview with one or more potential employers",
        "Complete career assessment",
        "Apply to one or more colleges/universities",
        "Accepted to one or more colleges/universities",
    }
)
GAR_BASELINE_GRACE_DAYS = 45
GAR_ENROLLMENT_START_KEYS = ("Enrollment Begin Date", "Enroll Start Date", "Enrollment BeginDate")
GAR_ENROLLMENT_END_KEYS = ("Enrollment Exit Date", "Enroll End Date", "Enrollment EndDate")

GarRecord: TypeAlias = dict[str, str | None | int]


def _gar_get_worksheet(workbook_path: Path, sheet_name: str):
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        available = ", ".join(workbook.sheetnames)
        workbook.close()
        raise ValueError(f"Worksheet not found: {sheet_name}. Available sheets: {available}")
    return workbook, workbook[sheet_name]


def _gar_validate_headers(headers: list[str | None], sheet_name: str) -> None:
    present = {header for header in headers if header}
    missing = sorted(header for header in GAR_REQUIRED_HEADERS if header not in present)
    if missing:
        joined = ", ".join(missing)
        raise ValueError(f"Worksheet {sheet_name} is missing required headers: {joined}")


def load_goal_progress_rows(
    workbook_path: Path,
    sheet_name: str = GOAL_PROGRESS_SHEET,
) -> list[GarRecord]:
    workbook, sheet = _gar_get_worksheet(workbook_path, sheet_name)
    rows = list(sheet.iter_rows(values_only=True))
    if len(rows) < GOAL_PROGRESS_HEADER_ROW:
        workbook.close()
        raise ValueError(f"Worksheet {sheet_name} does not contain header row {GOAL_PROGRESS_HEADER_ROW}")
    headers = [to_cell_text(value) for value in rows[GOAL_PROGRESS_HEADER_ROW - 1][:GAR_HEADER_COUNT]]
    _gar_validate_headers(headers, sheet_name)
    records: list[GarRecord] = []
    for row_number, row_values in enumerate(rows[GOAL_PROGRESS_HEADER_ROW:], start=GOAL_PROGRESS_HEADER_ROW + 1):
        record: GarRecord = {"row_number": row_number}
        for index, header in enumerate(headers):
            if header:
                record[header] = to_cell_text(row_values[index]) if index < len(row_values) else None
        records.append(record)
    workbook.close()
    return records


def _load_flexible_goal_progress_rows(workbook: Any) -> list[GarRecord]:
    worksheet = workbook[workbook.sheetnames[0]]
    header_row: list[str] | None = None
    header_index = -1
    for row_number, row in enumerate(
        worksheet.iter_rows(values_only=True, max_row=10),
        start=1,
    ):
        header_row = [to_cell_text(value) for value in row]
        normalized = [cell or "" for cell in header_row]
        if "Student ID" in normalized and "Goal Achievement" in normalized:
            header_index = row_number
            break
        header_row = None
    if header_row is None:
        return []
    records: list[GarRecord] = []
    for row_number, row_values in enumerate(
        worksheet.iter_rows(min_row=header_index + 1, values_only=True),
        start=header_index + 1,
    ):
        if not row_values or not any(row_values):
            continue
        record: GarRecord = {"row_number": row_number}
        for index, header in enumerate(header_row):
            if not header:
                continue
            value = row_values[index] if index < len(row_values) else None
            record[header] = to_cell_text(value)
        school = clean_value(str(record.get("School") or ""))
        if school and not clean_value(str(record.get("Home School") or "")):
            record["Home School"] = school
        records.append(record)
    return records


def load_goal_progress_gar_records(workbook_path: Path) -> list[GarRecord]:
    """Load rows for GAR from CIS_StudentProgress_Detail or archive variants."""
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if GOAL_PROGRESS_SHEET in workbook.sheetnames:
            workbook.close()
            try:
                return load_goal_progress_rows(workbook_path, sheet_name=GOAL_PROGRESS_SHEET)
            except ValueError:
                workbook = load_workbook(workbook_path, read_only=True, data_only=True)
        return _load_flexible_goal_progress_rows(workbook)
    finally:
        workbook.close()


def gar_percentage_to_grade(percentage: float) -> str | None:
    if percentage < 0 or percentage > 100:
        return None
    for threshold, grade in GAR_PERCENT_TO_GRADE_THRESHOLD:
        if percentage >= threshold:
            return grade
    return None


def gar_parse_date(value: object | None) -> date | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = clean_value(str(value))
    if text is None:
        return None
    if " " in text:
        text = text.split(" ", 1)[0].strip()
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y", "%d-%b-%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def gar_enrollment_start_date(record: GarRecord) -> date | None:
    for key in GAR_ENROLLMENT_START_KEYS:
        parsed = gar_parse_date(record.get(key))
        if parsed is not None:
            return parsed
    return None


def gar_enrollment_end_date(record: GarRecord) -> date | None:
    for key in GAR_ENROLLMENT_END_KEYS:
        parsed = gar_parse_date(record.get(key))
        if parsed is not None:
            return parsed
    return None


def gar_comparable_value(
    metric: str | None,
    value: str | None,
    context: dict[str, str] | None = None,
) -> float | None:
    metric_name = clean_value(metric) or ""
    grade = normalize_grade(value)
    if "Core Course Grades" in metric_name:
        if grade in GAR_GRADE_SCALE:
            return float(GAR_GRADE_SCALE[grade])
        if grade in GAR_ALTERNATE_GRADE_SCALE:
            return float(GAR_ALTERNATE_GRADE_SCALE[grade])
        if value is not None:
            try:
                pct = float(str(value).strip())
                if 0 <= pct <= 100:
                    letter = gar_percentage_to_grade(pct)
                    if letter:
                        return float(GAR_GRADE_SCALE[letter])
            except (ValueError, TypeError):
                pass
        return None
    numeric = parse_number(value)
    if numeric is not None:
        return numeric
    if grade in GAR_GRADE_SCALE:
        return float(GAR_GRADE_SCALE[grade])
    if grade in GAR_ALTERNATE_GRADE_SCALE:
        return float(GAR_ALTERNATE_GRADE_SCALE[grade])
    context_key = f"{metric_name}|{value}"
    if context is not None and context_key in context:
        scale_type = context[context_key]
        if scale_type == "alternate" and grade in GAR_ALTERNATE_GRADE_SCALE:
            return float(GAR_ALTERNATE_GRADE_SCALE[grade])
        if scale_type == "letter" and grade in GAR_GRADE_SCALE:
            return float(GAR_GRADE_SCALE[grade])
    if metric_name in GAR_YES_NO_METRICS:
        yes_no = clean_value(value)
        if yes_no is None:
            return None
        yes_no_norm = yes_no.lower()
        if yes_no_norm == "yes":
            return 1.0
        if yes_no_norm == "no":
            return 0.0
    if "engagement" in metric_name.lower():
        engagement_value = clean_value(value)
        if engagement_value is not None:
            engagement_norm = engagement_value.lower().strip()
            if "lower" in engagement_norm:
                return 1.0
            if "moderate" in engagement_norm:
                return 2.0
            if "higher" in engagement_norm or "high" in engagement_norm:
                return 3.0
    if "Reading Level" in metric_name and grade is not None:
        return parse_number(grade)
    if "Core Course Grades" in metric_name and value is not None:
        try:
            pct = float(str(value).strip())
            if 0 <= pct <= 100:
                letter = gar_percentage_to_grade(pct)
                if letter:
                    if context is not None:
                        context[context_key] = "letter"
                    return float(GAR_GRADE_SCALE[letter])
        except (ValueError, TypeError):
            pass
    return None


def gar_metric_direction(metric: str | None) -> str | None:
    metric_name = clean_value(metric)
    if metric_name is None:
        return None
    if "(SEAD)" in metric_name:
        return "higher"
    if "overall sead" in metric_name.lower() and "assessment score" in metric_name.lower():
        return "higher"
    if metric_name in GAR_METRIC_DIRECTIONS:
        return GAR_METRIC_DIRECTIONS[metric_name]
    lowered = metric_name.lower()
    if "grade" in lowered or "gpa" in lowered or "score" in lowered or "engagement" in lowered:
        return "higher"
    if "attendance rate" in lowered:
        return "higher" if "%" in metric_name else "lower"
    if "absence" in lowered or "suspension" in lowered or "tard" in lowered:
        return "lower"
    if "referral" in lowered or "incident" in lowered or "conduct" in lowered:
        return "lower"
    if "credits needed" in lowered or "credit completion" in lowered:
        return "higher"
    return None


def gar_baseline_grace_applies(record: GarRecord, metric: str | None, as_of: date) -> bool:
    baseline_missing = gar_comparable_value(metric, str(record.get("Baseline") or "")) is None
    if not baseline_missing:
        return False
    if gar_comparable_value(metric, str(record.get("Target") or "")) is None:
        return False
    if gar_comparable_value(metric, str(record.get("Achieved Value") or "")) is None:
        return False
    start = gar_enrollment_start_date(record)
    if start is None or start > as_of:
        return False
    return (as_of - start).days < GAR_BASELINE_GRACE_DAYS


def gar_baseline_grace_reason(record: GarRecord, as_of: date) -> str:
    start = gar_enrollment_start_date(record)
    end = gar_enrollment_end_date(record)
    if start is None:
        return "baseline grace period (enrollment start date missing)"
    days = (as_of - start).days
    end_note = f"enrollment end date {end.isoformat()}" if end else "enrollment end date not provided"
    return (
        f"baseline not audited yet: {days} days since enrollment begin ({start.isoformat()}); "
        f"{end_note} (within {GAR_BASELINE_GRACE_DAYS}-day window after enrollment start)"
    )


def gar_classify_row(
    record: GarRecord,
    context: dict[str, str] | None = None,
) -> dict[str, str | None]:
    metric = str(record.get("Metric") or "")
    existing = clean_value(str(record.get("Goal Achievement") or ""))
    if existing == GOAL_ACHIEVEMENT_EWGSPE:
        return {"expected": existing, "direction": gar_metric_direction(metric), "reason": "special_status"}
    direction = gar_metric_direction(metric)
    if context is None:
        context = {}
    baseline = gar_comparable_value(metric, str(record.get("Baseline") or ""), context)
    target = gar_comparable_value(metric, str(record.get("Target") or ""), context)
    achieved = gar_comparable_value(metric, str(record.get("Achieved Value") or ""), context)
    as_of = date.today()
    if direction is not None and gar_baseline_grace_applies(record, metric, as_of):
        return {"expected": None, "direction": direction, "reason": "baseline_grace_period"}
    if direction is None or None in {baseline, target, achieved}:
        return {"expected": None, "direction": direction, "reason": "manual_review"}
    meets_target = achieved >= target if direction == "higher" else achieved <= target
    improved = achieved > baseline if direction == "higher" else achieved < baseline
    if meets_target:
        expected = "Goal Met"
    elif improved:
        expected = "Goal Not Met, With Progress"
    else:
        expected = "Goal Not Met, No Progress"
    return {"expected": expected, "direction": direction, "reason": "rule_based"}


def gar_filter_records(
    records: list[GarRecord],
    school: str | None,
    metric: str | None,
) -> list[GarRecord]:
    filtered: list[GarRecord] = []
    for record in records:
        if school and clean_value(str(record.get("Home School") or "")) != school:
            continue
        if metric and clean_value(str(record.get("Metric") or "")) != metric:
            continue
        filtered.append(record)
    return filtered


def gar_summarize(
    records: list[GarRecord],
    context: dict[str, str] | None = None,
) -> dict[str, Any]:
    summary: Counter[str] = Counter()
    mismatches: list[dict[str, Any]] = []
    direction_counts: Counter[str] = Counter()
    if context is None:
        context = {}
    for record in records:
        classification = gar_classify_row(record, context)
        expected = classification["expected"]
        existing = clean_value(str(record.get("Goal Achievement") or ""))
        direction = classification["direction"]
        reason = classification["reason"]
        if direction:
            direction_counts[direction] += 1
        summary["total_rows"] += 1
        summary[f"reason::{reason}"] += 1
        if existing:
            summary[f"existing::{existing}"] += 1
        if expected:
            summary[f"expected::{expected}"] += 1
        if reason == "special_status":
            continue
        if expected and existing and expected == existing:
            summary["matched"] += 1
        elif expected and existing:
            summary["mismatched"] += 1
            mismatches.append(
                {
                    "row_number": record["row_number"],
                    "home_school": record.get("Home School"),
                    "goal": record.get("Goal"),
                    "metric": record.get("Metric"),
                    "baseline": record.get("Baseline"),
                    "target": record.get("Target"),
                    "achieved_value": record.get("Achieved Value"),
                    "existing_goal_achievement": existing,
                    "expected_goal_achievement": expected,
                    "direction": direction,
                }
            )
        else:
            summary["manual_review_rows"] += 1
    return {
        "summary": dict(summary),
        "direction_counts": dict(direction_counts),
        "mismatch_examples": mismatches[:25],
    }


def gar_explain_exception(
    record: GarRecord,
    classification: dict[str, str | None],
) -> dict[str, str]:
    reason = classification["reason"]
    expected = classification["expected"]
    existing = clean_value(str(record.get("Goal Achievement") or ""))
    metric = str(record.get("Metric") or "")
    if reason == "baseline_grace_period":
        return {
            "exception_type": "needs_manual_review",
            "reason": gar_baseline_grace_reason(record, date.today()),
        }
    if reason == "manual_review":
        if classification["direction"] is None:
            return {
                "exception_type": "needs_manual_review",
                "reason": "metric direction is not mapped",
            }
        missing_fields = []
        if gar_comparable_value(metric, str(record.get("Baseline") or "")) is None:
            missing_fields.append("baseline")
        if gar_comparable_value(metric, str(record.get("Target") or "")) is None:
            missing_fields.append("target")
        if gar_comparable_value(metric, str(record.get("Achieved Value") or "")) is None:
            missing_fields.append("achieved value")
        if missing_fields:
            if len(missing_fields) > 1:
                return {
                    "exception_type": "needs_manual_review",
                    "reason": "multiple values not comparable",
                }
            joined = ", ".join(missing_fields)
            return {
                "exception_type": "needs_manual_review",
                "reason": f"{joined} missing or not comparable",
            }
        return {
            "exception_type": "needs_manual_review",
            "reason": "needs manual review",
        }
    if reason == "special_status":
        return {
            "exception_type": "special_status",
            "reason": "special status preserved",
        }
    if expected and existing and expected != existing:
        return {
            "exception_type": "goal_achievement_mismatch",
            "reason": "recorded result does not match recalculated result",
        }
    return {
        "exception_type": "informational",
        "reason": "no exception",
    }


def gar_build_exception_rows(
    records: list[GarRecord],
    context: dict[str, str] | None = None,
) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    if context is None:
        context = {}
    for record in records:
        classification = gar_classify_row(record, context)
        expected = classification["expected"]
        existing = clean_value(str(record.get("Goal Achievement") or ""))
        direction = classification["direction"]
        reason = classification["reason"]
        if reason in ("manual_review", "baseline_grace_period"):
            explanation = gar_explain_exception(record, classification)
            rows.append(
                {
                    "exception_type": explanation["exception_type"],
                    "row_number": record["row_number"],
                    "home_school": record.get("Home School"),
                    "goal": record.get("Goal"),
                    "metric": record.get("Metric"),
                    "baseline": record.get("Baseline"),
                    "target": record.get("Target"),
                    "achieved_value": record.get("Achieved Value"),
                    "existing_goal_achievement": existing,
                    "expected_goal_achievement": expected,
                    "direction": direction,
                    "reason": explanation["reason"],
                }
            )
        elif expected and existing and expected != existing:
            explanation = gar_explain_exception(record, classification)
            rows.append(
                {
                    "exception_type": explanation["exception_type"],
                    "row_number": record["row_number"],
                    "home_school": record.get("Home School"),
                    "goal": record.get("Goal"),
                    "metric": record.get("Metric"),
                    "baseline": record.get("Baseline"),
                    "target": record.get("Target"),
                    "achieved_value": record.get("Achieved Value"),
                    "existing_goal_achievement": existing,
                    "expected_goal_achievement": expected,
                    "direction": direction,
                    "reason": explanation["reason"],
                }
            )
    return rows


def gar_build_report_details(
    records: list[GarRecord],
    context: dict[str, str] | None = None,
) -> dict[str, Any]:
    school_summary: dict[str, Counter[str]] = {}
    mismatch_schools: Counter[str] = Counter()
    mismatch_metrics: Counter[str] = Counter()
    mismatch_patterns: Counter[tuple[str, str]] = Counter()
    manual_review_schools: Counter[str] = Counter()
    manual_review_metrics: Counter[str] = Counter()
    if context is None:
        context = {}

    for record in records:
        school = clean_value(str(record.get("Home School") or "")) or "Unknown School"
        metric = clean_value(str(record.get("Metric") or "")) or "Unknown Metric"
        school_counts = school_summary.setdefault(school, Counter())
        school_counts["total_rows"] += 1

        classification = gar_classify_row(record, context)
        expected = classification["expected"]
        existing = clean_value(str(record.get("Goal Achievement") or ""))
        reason = classification["reason"]

        if reason == "special_status":
            school_counts["special_status"] += 1
        elif reason in ("manual_review", "baseline_grace_period"):
            school_counts["manual_review"] += 1
            manual_review_schools[school] += 1
            manual_review_metrics[metric] += 1
        elif expected and existing and expected == existing:
            school_counts["matched"] += 1
        elif expected and existing:
            school_counts["mismatched"] += 1
            mismatch_schools[school] += 1
            mismatch_metrics[metric] += 1
            mismatch_patterns[(existing, expected)] += 1

    school_rows = []
    for school in sorted(school_summary):
        counts = school_summary[school]
        school_rows.append(
            {
                "school": school,
                "total_rows": counts.get("total_rows", 0),
                "matched": counts.get("matched", 0),
                "mismatched": counts.get("mismatched", 0),
                "manual_review": counts.get("manual_review", 0),
                "special_status": counts.get("special_status", 0),
            }
        )

    return {
        "school_summary": school_rows,
        "top_mismatch_schools": mismatch_schools.most_common(10),
        "top_mismatch_metrics": mismatch_metrics.most_common(10),
        "top_mismatch_patterns": mismatch_patterns.most_common(10),
        "top_manual_review_schools": manual_review_schools.most_common(10),
        "top_manual_review_metrics": manual_review_metrics.most_common(10),
    }


def gar_write_markdown_report(payload: dict[str, Any], report_path: Path) -> None:
    summary = payload["summary"]
    direction_counts = payload["direction_counts"]
    mismatch_examples = payload["mismatch_examples"]
    report_details = payload.get("report_details", {})
    report_path.parent.mkdir(parents=True, exist_ok=True)
    lines = [
        "# Goal Achievement Audit Report",
        "",
        f"- Workbook: `{payload['workbook']}`",
        f"- Worksheet: `{payload['sheet_name']}`",
        f"- Rows analyzed: {payload['record_count']}",
        f"- Matched: {summary.get('matched', 0)}",
        f"- Mismatched: {summary.get('mismatched', 0)}",
        f"- Manual review: {summary.get('manual_review_rows', 0)}",
        "",
        "## Direction counts",
    ]
    for key, value in sorted(direction_counts.items()):
        lines.append(f"- `{key}`: {value}")
    top_mismatch_schools = report_details.get("top_mismatch_schools", [])
    if top_mismatch_schools:
        lines.extend(["", "## Top mismatch schools"])
        for school, count in top_mismatch_schools:
            lines.append(f"- `{school}`: {count}")
    top_mismatch_metrics = report_details.get("top_mismatch_metrics", [])
    if top_mismatch_metrics:
        lines.extend(["", "## Top mismatch metrics"])
        for metric, count in top_mismatch_metrics:
            lines.append(f"- `{metric}`: {count}")
    top_manual_review_metrics = report_details.get("top_manual_review_metrics", [])
    if top_manual_review_metrics:
        lines.extend(["", "## Top manual review metrics"])
        for metric, count in top_manual_review_metrics:
            lines.append(f"- `{metric}`: {count}")
    top_mismatch_patterns = report_details.get("top_mismatch_patterns", [])
    if top_mismatch_patterns:
        lines.extend(["", "## Mismatch patterns"])
        for (existing, expected), count in top_mismatch_patterns:
            lines.append(f"- recorded `{existing}` -> expected `{expected}`: {count}")
    school_summary = report_details.get("school_summary", [])
    if school_summary:
        lines.extend(
            [
                "",
                "## School summary",
                "",
                "| School | Rows | Matched | Mismatched | Manual review | Special status |",
                "|---|---:|---:|---:|---:|---:|",
            ]
        )
        for item in school_summary:
            lines.append(
                f"| {item['school']} | {item['total_rows']} | {item['matched']} | "
                f"{item['mismatched']} | {item['manual_review']} | {item['special_status']} |"
            )
    if mismatch_examples:
        lines.extend(["", "## Sample mismatches"])
        for item in mismatch_examples[:10]:
            lines.append(
                "- "
                f"{item['home_school']} | {item['metric']} | "
                f"recorded `{item['existing_goal_achievement']}` vs expected `{item['expected_goal_achievement']}`"
            )
    report_path.write_text("\n".join(lines) + "\n")


def _empty_gar_outcome_counts() -> dict[str, int]:
    return {label: 0 for label in GAR_OUTCOME_LABELS}


def _empty_gar_outcome_bucket() -> dict[str, Any]:
    return {
        "recorded": _empty_gar_outcome_counts(),
        "expected": _empty_gar_outcome_counts(),
        "matched": 0,
        "mismatched": 0,
        "manual_review": 0,
        "special_status": 0,
        "total_rows": 0,
    }


def _finalize_gar_outcome_bucket(bucket: dict[str, Any]) -> dict[str, Any]:
    classifiable = int(bucket["matched"]) + int(bucket["mismatched"])
    recorded_total = sum(int(bucket["recorded"][label]) for label in GAR_OUTCOME_LABELS)
    expected_total = sum(int(bucket["expected"][label]) for label in GAR_OUTCOME_LABELS)
    mismatch_rate_pct: float | None = None
    if classifiable:
        mismatch_rate_pct = round(100 * int(bucket["mismatched"]) / classifiable, 1)
    recorded_goal_met_pct: float | None = None
    if recorded_total:
        recorded_goal_met_pct = round(
            100 * int(bucket["recorded"]["Goal Met"]) / recorded_total,
            1,
        )
    expected_goal_met_pct: float | None = None
    if expected_total:
        expected_goal_met_pct = round(
            100 * int(bucket["expected"]["Goal Met"]) / expected_total,
            1,
        )
    recorded_not_met = sum(
        int(bucket["recorded"][label]) for label in GAR_GOAL_NOT_MET_LABELS
    )
    expected_not_met = sum(
        int(bucket["expected"][label]) for label in GAR_GOAL_NOT_MET_LABELS
    )
    recorded_goal_not_met_pct: float | None = None
    if recorded_total:
        recorded_goal_not_met_pct = round(100 * recorded_not_met / recorded_total, 1)
    expected_goal_not_met_pct: float | None = None
    if expected_total:
        expected_goal_not_met_pct = round(100 * expected_not_met / expected_total, 1)
    return {
        **bucket,
        "classifiable_rows": classifiable,
        "eligible_rows": classifiable,
        "mismatch_rate_pct": mismatch_rate_pct,
        "recorded_goal_met_pct": recorded_goal_met_pct,
        "expected_goal_met_pct": expected_goal_met_pct,
        "recorded_goal_not_met_pct": recorded_goal_not_met_pct,
        "expected_goal_not_met_pct": expected_goal_not_met_pct,
    }


def gar_outcome_rollup(
    records: list[GarRecord],
    context: dict[str, str] | None = None,
) -> dict[str, Any]:
    if context is None:
        context = {}
    global_bucket = _empty_gar_outcome_bucket()
    schools: dict[str, dict[str, Any]] = {}

    def school_bucket(school: str) -> dict[str, Any]:
        if school not in schools:
            schools[school] = _empty_gar_outcome_bucket()
        return schools[school]

    for record in records:
        school = clean_value(str(record.get("Home School") or "")) or "(unknown)"
        school_counts = school_bucket(school)
        for bucket in (global_bucket, school_counts):
            bucket["total_rows"] += 1

        classification = gar_classify_row(record, context)
        expected = classification["expected"]
        existing = clean_value(str(record.get("Goal Achievement") or ""))
        reason = classification["reason"]

        if reason == "special_status":
            for bucket in (global_bucket, school_counts):
                bucket["special_status"] += 1
            continue
        if reason in ("manual_review", "baseline_grace_period"):
            for bucket in (global_bucket, school_counts):
                bucket["manual_review"] += 1
            continue

        if existing in GAR_OUTCOME_LABELS:
            for bucket in (global_bucket, school_counts):
                bucket["recorded"][existing] += 1
        if expected in GAR_OUTCOME_LABELS:
            for bucket in (global_bucket, school_counts):
                bucket["expected"][expected] += 1

        if expected and existing:
            if expected == existing:
                for bucket in (global_bucket, school_counts):
                    bucket["matched"] += 1
            else:
                for bucket in (global_bucket, school_counts):
                    bucket["mismatched"] += 1

    return {
        "global": _finalize_gar_outcome_bucket(global_bucket),
        "schools": {
            school: _finalize_gar_outcome_bucket(counts)
            for school, counts in sorted(schools.items())
        },
    }


def gar_write_outcome_rollup_json(rollup: dict[str, Any], report_path: Path) -> None:
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text(json.dumps(rollup, indent=2, sort_keys=True) + "\n", encoding="utf-8")


def gar_write_csv_report(exception_rows: list[dict[str, Any]], report_path: Path) -> None:
    report_path.parent.mkdir(parents=True, exist_ok=True)
    fieldnames = [
        "exception_type",
        "row_number",
        "home_school",
        "goal",
        "metric",
        "baseline",
        "target",
        "achieved_value",
        "existing_goal_achievement",
        "expected_goal_achievement",
        "direction",
        "reason",
    ]
    with report_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(exception_rows)


def run_gar_audit(
    records: list[GarRecord],
    *,
    school: str | None = None,
    metric: str | None = None,
) -> dict[str, Any]:
    context: dict[str, str] = {}
    filtered = gar_filter_records(records, school, metric)
    results = gar_summarize(filtered, context)
    exception_rows = gar_build_exception_rows(filtered, context)
    report_details = gar_build_report_details(filtered, context)
    outcome_rollup = gar_outcome_rollup(filtered, context)
    return {
        "exception_rows": exception_rows,
        "summary": results["summary"],
        "direction_counts": results["direction_counts"],
        "mismatch_examples": results["mismatch_examples"],
        "report_details": report_details,
        "outcome_rollup": outcome_rollup,
        "record_count": len(filtered),
        "sheet_name": GOAL_PROGRESS_SHEET,
    }


if __name__ == "__main__":
    raise SystemExit(main())