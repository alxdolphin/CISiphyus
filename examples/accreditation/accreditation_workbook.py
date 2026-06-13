#!/usr/bin/env python3
"""Parse CISDM accreditation .xlsx exports into site-level row dicts."""

from __future__ import annotations

from itertools import chain
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

REQUIRED_SHEETS = (
    "Accreditation Site Coordination",
    "Accreditation Tier I",
    "Accreditation Case Management",
)


def to_cell_text(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def normalize_header(value: object) -> str:
    if value is None:
        return ""
    return " ".join(str(value).replace("\n", " ").split())


def is_coordination_header_row(row_values: tuple[Any, ...]) -> bool:
    if not row_values or len(row_values) < 2:
        return False
    first = to_cell_text(row_values[0])
    second = to_cell_text(row_values[1])
    return first == "Organization" and second == "School"


def disambiguate_headers(labels: list[str]) -> list[str]:
    counts: dict[str, int] = {}
    out: list[str] = []
    for label in labels:
        n = counts.get(label, 0) + 1
        counts[label] = n
        out.append(label if n == 1 else f"{label}__{n}")
    return out


def find_column_index(headers: list[str], *substrings: str) -> int | None:
    """Return first column index whose normalized header contains all substrings."""
    return find_column_index_excluding(headers, *substrings)


def find_column_index_excluding(
    headers: list[str],
    *substrings: str,
    exclude: tuple[str, ...] = (),
) -> int | None:
    needles = [s.lower() for s in substrings]
    banned = [s.lower() for s in exclude]
    for index, header in enumerate(headers):
        normalized = normalize_header(header).lower()
        if not normalized:
            continue
        if any(ban in normalized for ban in banned):
            continue
        if all(needle in normalized for needle in needles):
            return index
    return None


def require_column(headers: list[str], sheet_name: str, *substrings: str) -> int:
    index = find_column_index(headers, *substrings)
    if index is not None:
        return index
    available = [normalize_header(h) for h in headers if h]
    raise ValueError(
        f"Worksheet {sheet_name!r} missing column matching {substrings!r}. "
        f"Available headers: {available}"
    )


def parse_number(value: object) -> tuple[float | None, str | None]:
    text = to_cell_text(value)
    if text is None:
        return None, None
    cleaned = text.replace(",", "").replace("%", "").strip()
    try:
        return float(cleaned), None
    except ValueError:
        return None, f"non_numeric:{text!r}"


def parse_int(value: object) -> tuple[int, str | None]:
    number, warning = parse_number(value)
    if number is None:
        return 0, warning
    return int(number), warning


def parse_sheet_rows(workbook_path: Path, sheet_name: str) -> list[dict[str, Any]]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            available = ", ".join(workbook.sheetnames)
            raise ValueError(f"Worksheet not found: {sheet_name}. Available: {available}")

        sheet = workbook[sheet_name]
        row_iter: Iterable[tuple[Any, ...]] = sheet.iter_rows(values_only=True)
        iterator = iter(row_iter)
        row1 = next(iterator, None)
        if row1 is None:
            return []

        row2 = next(iterator, None)
        if row2 is not None and is_coordination_header_row(row2):
            header_cells = row2
            data_iterator: Iterable[tuple[Any, ...]] = iterator
        else:
            header_cells = row1
            data_iterator = chain([row2], iterator) if row2 is not None else iterator

        raw_headers = [normalize_header(value) for value in header_cells]
        label_indices: list[tuple[int, str]] = [
            (index, header) for index, header in enumerate(raw_headers) if header
        ]
        if not label_indices:
            return []

        labels = [pair[1] for pair in label_indices]
        unique_labels = disambiguate_headers(labels)
        columns = [(label_indices[i][0], unique_labels[i]) for i in range(len(label_indices))]

        records: list[dict[str, Any]] = []
        for excel_row, row_values in enumerate(data_iterator, start=1):
            record: dict[str, Any] = {"_excel_row": excel_row}
            has_data = False
            for col_index, key in columns:
                cell = row_values[col_index] if col_index < len(row_values) else None
                text = to_cell_text(cell)
                record[key] = text
                if text is not None:
                    has_data = True
            if has_data:
                records.append(record)
        return records
    finally:
        workbook.close()


def load_accreditation_workbook(workbook_path: Path) -> dict[str, list[dict[str, Any]]]:
    workbook_path = workbook_path.resolve()
    return {
        sheet_name: parse_sheet_rows(workbook_path, sheet_name)
        for sheet_name in REQUIRED_SHEETS
    }
