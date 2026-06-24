#!/usr/bin/env python3
"""Synthetic student metrics + accreditation drilldown workbooks for cross-audit tests."""

from __future__ import annotations

from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook

GOAL_PROGRESS_SHEET = "CIS_StudentProgress_Detail"
GOAL_PROGRESS_HEADERS = [
    "Student ID",
    "Home School",
    "Goal",
    "Metric",
    "Baseline",
    "Target",
    "Achieved Value",
    "Goal Achievement",
    "Enrollment Begin Date",
    "Enrollment Exit Date",
]

METRICS_HEADERS = [
    "Organization",
    "School",
    "Student ID",
    "Client ID",
    "Student Name",
    "Grade Level",
    "Case Manager",
    "Goal",
    "Metric",
    "Baseline",
    "Target",
    "Latest Progress",
]

DRILLDOWN_HEADERS = [
    "Organization",
    "School",
    "Student ID",
    "Client ID",
    "Student Name",
    "# ABCS Metrics Assigned",
    "# Non-ABCS Metrics Assigned",
    "# of Metrics with Goal Achievement entered",
    "Goal Achievement Entered for ALL assigned goals?",
]


def _write_coordination_style_sheet(
    workbook: Workbook,
    sheet_name: str,
    headers: list[str],
    rows: list[list[object]],
) -> None:
    sheet = workbook.create_sheet(sheet_name)
    sheet.cell(row=1, column=1, value=f"{sheet_name} banner")
    for index, header in enumerate(headers, start=1):
        sheet.cell(row=2, column=index, value=header)
    for row_index, row in enumerate(rows, start=3):
        for col_index, value in enumerate(row, start=1):
            sheet.cell(row=row_index, column=col_index, value=value)


def build_metrics_fixture(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    headers = [
        "Organization",
        "School",
        "Student ID",
        "Client ID",
        "Goal",
        "Metric",
        "Baseline",
        "Target",
        "Enrollment Status",
        "School Year",
        "Case Manager",
        "Latest Progress",
        "Goal Achievement",
    ]
    for index, header in enumerate(headers, start=1):
        sheet.cell(row=1, column=index, value=header)
    rows = [
        [
            "CIS of Eastern Pennsylvania",
            "Test School",
            "1001",
            "2001",
            "Improve Attendance",
            "Attendance Rate (%)",
            "85",
            "90",
            "Active",
            "2025-26",
            "Coordinator",
            "91",
            "Goal Not Met, No Progress",
        ],
        [
            "CIS of Eastern Pennsylvania",
            "Test School",
            "1002",
            "2002",
            "Improve Attendance",
            "Attendance Rate (%)",
            "80",
            "88",
            "Active",
            "2025-26",
            "Coordinator",
            "88",
            None,
        ],
        [
            "CIS of Eastern Pennsylvania",
            "Test School",
            "1003",
            "2003",
            "Improve Attendance",
            "Attendance Rate (%)",
            "",
            "",
            "Active",
            "2025-26",
            "Coordinator",
            "85",
            None,
        ],
    ]
    for row_index, row in enumerate(rows, start=2):
        for col_index, value in enumerate(row, start=1):
            sheet.cell(row=row_index, column=col_index, value=value)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def build_goal_progress_fixture(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = GOAL_PROGRESS_SHEET
    sheet.cell(row=1, column=1, value="Goal Tracking export")
    sheet.cell(row=2, column=1, value="Student progress detail")
    for index, header in enumerate(GOAL_PROGRESS_HEADERS, start=1):
        sheet.cell(row=3, column=index, value=header)

    recent_start = (date.today() - timedelta(days=10)).isoformat()
    rows = [
        [
            "1001",
            "Test School",
            "Improve Attendance",
            "Attendance Rate (%)",
            "85",
            "90",
            "91",
            "Goal Not Met, No Progress",
            "2025-08-01",
            None,
        ],
        [
            "1002",
            "Test School",
            "Improve Attendance",
            "Attendance Rate (%)",
            None,
            "88",
            "87",
            "Goal Not Met, No Progress",
            recent_start,
            None,
        ],
        [
            "1003",
            "Test School",
            "Improve Academics",
            "Custom Widget Score",
            "1",
            "2",
            "3",
            "Goal Met",
            "2025-08-01",
            None,
        ],
    ]
    for row_index, row in enumerate(rows, start=4):
        for col_index, value in enumerate(row, start=1):
            sheet.cell(row=row_index, column=col_index, value=value)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def build_accreditation_fixture(path: Path) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    _write_coordination_style_sheet(
        workbook,
        "Accreditation Student Drilldown",
        DRILLDOWN_HEADERS,
        [
            [
                "CIS of Eastern Pennsylvania",
                "Test School",
                "1001",
                "2001",
                "Student One",
                1,
                0,
                1,
                "Yes",
            ],
            [
                "CIS of Eastern Pennsylvania",
                "Test School",
                "1002",
                "2002",
                "Student Two",
                1,
                0,
                0,
                "No",
            ],
        ],
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def build_fixture(path: Path) -> None:
    """Backward-compatible alias: metrics-only fixture."""
    build_metrics_fixture(path)
